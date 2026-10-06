# Fase 4: exportar tus datos a Excel, probar una regla antes de guardarla, fusionar y ocultar categorías. Datos inventados.
import base64, io, os, shutil, sys, tempfile, unittest
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from financebuddy import plantilla, servidor

class Base(unittest.TestCase):
    def setUp(self):
        self.dir = tempfile.mkdtemp()
        self.app = servidor.App(self.dir, hoy="2026-09-30")
        self.alm = self.app.alm
    def tearDown(self):
        self.alm.cerrar(); shutil.rmtree(self.dir, ignore_errors=True)
    def api(self, ruta, d=None): return self.app.manejar(ruta, d or {})
    def mov(self, texto, imp, cat, clase="gasto", fecha="2026-09-10"):
        return self.alm.guardar("movimiento", {"fecha": fecha, "clase": clase, "categoria": cat, "importe": imp, "concepto": texto, "cuenta": "Nómina", "ext_texto": texto})
    def preparar(self):
        for n, g in (("Comer fuera", "variable"), ("Supermercado", "variable"), ("Ocio", "variable"), ("Sueldo", "ingreso")):
            if not self.alm.buscar("categoria", nombre=n): self.alm.guardar("categoria", {"nombre": n, "grupo": g})

class TestExportar(Base):
    def test_excel_con_una_hoja_por_cosa_y_sin_formulas(self):
        from openpyxl import load_workbook
        self.preparar()
        self.mov("=HYPERLINK(\"http://x\")", 12.5, "Ocio")
        self.mov("Nómina empresa", 1500, "Sueldo", "ingreso")
        r = self.api("/api/exportar_datos")
        self.assertTrue(r["ok"] and r["nombre"].endswith(".xlsx"))
        wb = load_workbook(io.BytesIO(base64.b64decode(r["contenido"])))
        for hoja in ("Léeme", "Movimientos", "Cuentas", "Categorías", "Reglas", "Inversión", "Aportaciones", "Fijos"):
            self.assertIn(hoja, wb.sheetnames)
        ws = wb["Movimientos"]
        cab = [c.value for c in ws[1]]
        self.assertEqual(cab[:5], ["Fecha", "Tipo", "Concepto", "Categoría", "Importe"])
        filas = {f[2].value: f for f in ws.iter_rows(min_row=2)}
        f = filas['=HYPERLINK("http://x")']
        self.assertEqual(f[2].data_type, "s")  # texto, no fórmula
        self.assertEqual((f[4].value, f[5].value), (12.5, -12.5))
        self.assertEqual(filas["Nómina empresa"][5].value, 1500)
        self.assertEqual(wb["Movimientos"].max_row - 1, len(self.alm.todos("movimiento")))

    def test_vacio_tambien_se_exporta(self):
        from openpyxl import load_workbook
        r = self.api("/api/exportar_datos")
        wb = load_workbook(io.BytesIO(base64.b64decode(r["contenido"])))
        self.assertEqual(wb["Movimientos"].max_row, 1)

class TestProbarRegla(Base):
    def setUp(self):
        super().setUp(); self.preparar()
        self.a = self.mov("Compra Mercadona Valencia", 30, "Ocio")
        self.b = self.mov("Compra Mercadona Madrid", 20, "Supermercado")
        self.c = self.mov("Compra Zara", 50, "Ocio")
        self.d = self.mov("Devolución Mercadona", 10, "Ocio", "reembolso")
        self.alm.guardar("movimiento", {"fecha": "2026-09-11", "clase": "gasto", "categoria": "Ocio", "importe": 5, "concepto": "Mercadona a mano"})  # sin extracto: no cuenta

    def test_cuenta_lo_que_casa_y_lo_que_cambiaria(self):
        r = self.api("/api/regla/probar", {"patron": "mercadona", "clase": "gasto", "categoria": "Supermercado"})
        self.assertEqual((r["n"], r["cambian"], r["ya"]), (2, 1, 1))  # el reembolso (entra) no se mezcla con un gasto
        self.assertTrue(r["aplicable"]); self.assertEqual(r["muestra"][0]["id"], self.a)
        self.assertLessEqual(len(r["muestra"]), 8)

    def test_probar_no_toca_nada_y_aplicar_si(self):
        antes = self.alm.todos("movimiento")
        self.api("/api/regla/probar", {"patron": "mercadona", "clase": "gasto", "categoria": "Supermercado"})
        self.assertEqual(self.alm.todos("movimiento"), antes)
        r = self.api("/api/regla/aplicar", {"patron": "mercadona", "clase": "gasto", "categoria": "Supermercado"})
        self.assertIn("1 movimiento", r["mensaje"])
        self.assertEqual(self.alm.obtener("movimiento", self.a)["categoria"], "Supermercado")
        self.assertEqual(self.alm.obtener("movimiento", self.c)["categoria"], "Ocio")      # otro comercio, intacto
        self.assertEqual(self.alm.obtener("movimiento", self.d)["categoria"], "Ocio")      # reembolso, intacto
        self.assertEqual(self.api("/api/regla/probar", {"patron": "mercadona", "clase": "gasto", "categoria": "Supermercado"})["cambian"], 0)
        # y se puede deshacer
        self.assertTrue(self.api("/api/deshacer")["ok"])
        self.assertEqual(self.alm.obtener("movimiento", self.a)["categoria"], "Ocio")

    def test_patron_corto_o_transferencia_no_se_aplican(self):
        self.assertTrue(self.api("/api/regla/probar", {"patron": "m", "clase": "gasto", "categoria": "Ocio"})["corto"])
        r = self.api("/api/regla/probar", {"patron": "mercadona", "clase": "transferencia"})
        self.assertFalse(r["aplicable"]); self.assertGreater(r["n"], 0)
        with self.assertRaises(ValueError): self.api("/api/regla/aplicar", {"patron": "mercadona", "clase": "transferencia", "categoria": "Ocio"})
        with self.assertRaises(ValueError): self.api("/api/regla/aplicar", {"patron": "mercadona", "clase": "gasto", "categoria": "No existe"})

class TestCategorias(Base):
    def setUp(self):
        super().setUp(); self.preparar()
        self.alm.guardar("categoria", {"nombre": "Cenas", "grupo": "variable", "presupuesto": 50})
        self.alm.guardar("categoria", {"nombre": "Comer fuera", "grupo": "variable", "presupuesto": 100}, self.alm.buscar("categoria", nombre="Comer fuera")[0]["id"])
        self.m = self.mov("Restaurante", 40, "Cenas")
        self.alm.guardar("recurrente", {"nombre": "Cena mensual", "clase": "gasto", "categoria": "Cenas", "importe": 40, "desde": "2026-01-01"})
        self.alm.guardar("regla", {"patron": "restaurante", "categoria": "Cenas", "clase": "gasto", "origen": "usuario"})

    def test_vista_previa_no_cambia_nada(self):
        v = self.api("/api/categoria/fusionar", {"origen": "Cenas", "destino": "Comer fuera", "previa": True})
        self.assertEqual((v["movimientos"], v["fijos"], v["reglas"], v["presupuesto"]), (1, 1, 1, 150))
        self.assertEqual(self.alm.obtener("movimiento", self.m)["categoria"], "Cenas")
        self.assertTrue(self.alm.buscar("categoria", nombre="Cenas"))

    def test_fusionar_reasigna_todo_y_se_deshace(self):
        r = self.api("/api/categoria/fusionar", {"origen": "Cenas", "destino": "Comer fuera"})
        self.assertTrue(r["ok"])
        self.assertEqual(self.alm.obtener("movimiento", self.m)["categoria"], "Comer fuera")
        self.assertEqual(self.alm.todos("recurrente")[0]["categoria"], "Comer fuera")
        self.assertEqual(self.alm.buscar("regla", patron="restaurante")[0]["categoria"], "Comer fuera")
        self.assertFalse(self.alm.buscar("categoria", nombre="Cenas"))
        self.assertEqual(self.alm.buscar("categoria", nombre="Comer fuera")[0]["presupuesto"], 150)
        self.assertEqual(self.alm.config("categorias_fusionadas"), {"Cenas": "Comer fuera"})
        self.assertTrue(self.api("/api/deshacer")["ok"])
        self.assertTrue(self.alm.buscar("categoria", nombre="Cenas"))
        self.assertEqual(self.alm.obtener("movimiento", self.m)["categoria"], "Cenas")
        self.assertEqual(self.alm.buscar("categoria", nombre="Comer fuera")[0]["presupuesto"], 100)

    def test_no_mezcla_ingresos_con_gastos_ni_la_misma(self):
        with self.assertRaises(ValueError): self.api("/api/categoria/fusionar", {"origen": "Cenas", "destino": "Sueldo"})
        with self.assertRaises(ValueError): self.api("/api/categoria/fusionar", {"origen": "Cenas", "destino": "Cenas"})
        with self.assertRaises(ValueError): self.api("/api/categoria/fusionar", {"origen": "Cenas", "destino": "Nada"})
        self.assertTrue(self.alm.buscar("categoria", nombre="Cenas"))

    def test_ocultar_conserva_los_movimientos_y_se_puede_mostrar(self):
        self.api("/api/categoria/ocultar", {"nombres": ["Cenas"]})
        c = self.alm.buscar("categoria", nombre="Cenas")[0]
        self.assertTrue(c["oculta"]); self.assertEqual(self.alm.obtener("movimiento", self.m)["categoria"], "Cenas")
        from financebuddy import jev
        self.assertNotIn("Cenas", [n for n, _ in jev.categorias_de(self.alm)[0]])
        self.api("/api/categoria/ocultar", {"nombres": ["Cenas"], "ocultar": False})
        self.assertFalse(self.alm.buscar("categoria", nombre="Cenas")[0]["oculta"])

    def test_uso_cuenta_movimientos_fijos_y_reglas(self):
        u = self.api("/api/categoria/uso")["uso"]
        self.assertEqual(u["Cenas"], {"movimientos": 1, "fijos": 1, "reglas": 1})
        self.alm.guardar("categoria", {"nombre": "Vacía", "grupo": "variable"})
        self.assertEqual(self.api("/api/categoria/uso")["uso"]["Vacía"], {"movimientos": 0, "fijos": 0, "reglas": 0})

    def test_la_plantilla_no_recrea_una_categoria_fusionada(self):
        nombre = plantilla.CATEGORIAS[0][0]
        for c in self.alm.buscar("categoria", nombre=nombre): self.alm.borrar("categoria", c["id"])
        self.alm.set_config("categorias_fusionadas", {nombre: "Ocio"}); self.alm.set_config("plantilla_version", 1)
        plantilla.instalar(self.alm)
        self.assertFalse(self.alm.buscar("categoria", nombre=nombre))

if __name__ == "__main__":
    unittest.main()
