# Pruebas de la importación y la base de datos (sin navegador). Uso: python -m unittest pruebas.test_importar -v
import datetime, io, os, shutil, sys, tempfile, unittest
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from financebuddy import importar as IM, lectura as L, plantilla, rutas
from financebuddy.almacen import Almacen

def escribir(ruta, enc, texto):
    with io.open(ruta, "w", encoding=enc) as fh: fh.write(texto)

def excel_santander(ruta, filas, saldo_inicial=1000.0, romper=False, iban="ES00 0000 0000 0000 0000", titular=None):
    """filas: [(fecha ISO, concepto, importe)] en orden cronológico → Excel con el formato de Santander (más reciente arriba)."""
    import openpyxl
    wb = openpyxl.Workbook(); ws = wb.active
    ws.append([None, None, "Cuenta", "Fecha"]); ws.append([None, None, iban, "30/09/2026 | 10:00:00"])
    ws.append([None, None, "Titular", "Saldo"]); ws.append([None, None, titular or "", "1.000,00€ EUR"])
    ws.append(["Movimientos de la cuenta"]); ws.append([])
    ws.append(["Fecha operación", "Fecha valor", "Concepto", "Importe", "Saldo", "Divisa"])
    s, out = saldo_inicial, []
    for i, (f, c, imp) in enumerate(filas):
        s = round(s + imp, 2)
        out.append([datetime.datetime.fromisoformat(f), datetime.datetime.fromisoformat(f), c, imp, s + (5 if romper and i == 2 else 0), "EUR"])
    for r in reversed(out): ws.append(r)
    wb.save(ruta)

FILAS = [
    ("2026-09-01", "Compra Mercadona, Madrid, Tarj. :*1234", -45.20),
    ("2026-09-02", "Bizum a favor de Ana Ruiz concepto cena", -30.00),
    ("2026-09-03", "Traspaso: Ahorro", -200.00),
    ("2026-09-05", "Pago Movil En Ferreteria Lopez, Madrid", -12.50),
    ("2026-09-06", "Compra Lidl, Madrid", -10.00),
    ("2026-09-06", "Compra Lidl, Madrid", -10.00),  # dos compras idénticas el mismo día
    ("2026-09-10", "Transferencia Nomina Empresa SL", 1800.00),
    ("2026-09-11", "Transferencia de Juan Perez", 50.00),
]

class Base(unittest.TestCase):
    def setUp(self):
        self.dir = tempfile.mkdtemp()
        self.c = rutas.Carpeta(self.dir)
        self.a = Almacen(self.c.db)
        plantilla.instalar(self.a)
        self.a.guardar("cuenta", {"nombre": "Nómina", "tipo": "corriente", "extracto": True})
        self.a.guardar("cuenta", {"nombre": "Ahorro", "tipo": "ahorro"})
        self.a.guardar("cuenta", {"nombre": "Bróker", "tipo": "broker"})
        self.a.guardar("activo", {"nombre": "Fondo MSCI", "clase": "fondo", "cuenta": "Bróker", "patrones": ["msci world"]})
    def tearDown(self):
        self.a.cerrar(); shutil.rmtree(self.dir, ignore_errors=True)
    def extracto(self, filas=FILAS, nombre="movs.xlsx", **kw):
        ruta = os.path.join(self.c.banco, nombre); excel_santander(ruta, filas, **kw); return ruta
    def importar(self, ruta, tipo="banco", cuenta="Nómina"):
        return IM.importar_archivo(self.a, self.c, ruta, tipo, cuenta)
    def movs(self): return self.a.todos("movimiento")

class TestBanco(Base):
    def test_pide_la_cuenta_la_primera_vez(self):
        with self.assertRaises(IM.NecesitaCuenta): IM.importar_archivo(self.a, self.c, self.extracto(), "banco")
        r = self.importar(self.extracto())
        self.assertTrue(r["ok"])
        self.assertEqual(next(p for p in self.a.todos("perfil") if p["nombre"] == "Santander")["cuenta"], "Nómina")
        r2 = IM.importar_archivo(self.a, self.c, self.extracto(nombre="otro.xlsx"), "banco")  # ya no pregunta
        self.assertTrue(r2["ok"])

    def test_clasifica(self):
        r = self.importar(self.extracto())
        self.assertEqual((r["nuevas"], r["dudas"]), (6, 2))
        por = {m["ext_texto"]: m for m in self.movs()}
        self.assertEqual(por["Compra Mercadona, Madrid, Tarj. :*1234"]["categoria"], "Supermercado")
        self.assertEqual(por["Bizum a favor de Ana Ruiz concepto cena"]["categoria"], "Comer fuera")
        self.assertEqual((por["Traspaso: Ahorro"]["clase"], por["Traspaso: Ahorro"].get("destino")), ("transferencia", "Ahorro"))
        self.assertEqual((por["Transferencia Nomina Empresa SL"]["clase"], por["Transferencia Nomina Empresa SL"]["categoria"]), ("ingreso", "Nómina"))
        self.assertEqual(sum(1 for m in self.movs() if m["ext_texto"] == "Compra Lidl, Madrid"), 2)
        self.assertEqual(len(self.a.todos("pendiente")), 2)
        self.assertFalse(os.path.exists(os.path.join(self.c.banco, "movs.xlsx")))  # movido a Procesados
        self.assertEqual(self.a.config("saldo_extracto:Nómina"), {"fecha": "2026-09-11", "saldo": 2542.3})

    def test_reglas_de_serie(self):
        from financebuddy import clasificar as C
        reglas = C.ordenar_reglas(self.a.todos("regla"))
        cat = lambda t: (C.regla_para(t, reglas) or {}).get("categoria")
        self.assertEqual(cat("PAGO MOVIL EN DIA 25013, MADRID ES, TARJ. :*001234"), "Supermercado")
        self.assertEqual(cat("PAGO MOVIL EN PLENERGY US 58, MALAGA ES"), "Coche")
        self.assertEqual(cat("PAGO MOVIL EN GAME 458 MADRID, MADRID ES"), "Compras")
        self.assertIsNone(cat("Bizum a favor de Ana concepto dia de playa"))

    def test_reimportar_no_duplica(self):
        self.importar(self.extracto())
        r = self.importar(self.extracto())
        self.assertEqual((r["nuevas"], r["dudas"], r["existentes"]), (0, 0, 8))
        self.assertEqual((len(self.movs()), len(self.a.todos("pendiente"))), (6, 2))

    def test_resolver_y_aprender(self):
        self.importar(self.extracto())
        p = next(x for x in self.a.todos("pendiente") if "Ferreteria" in x["fila"]["texto"])
        IM.resolver(self.a, p["id"], {"accion": "guardar", "clase": "gasto", "categoria": "Hogar", "concepto": "Ferretería", "recordar": True, "patron": "ferreteria lopez"})
        q = next(x for x in self.a.todos("pendiente") if "Juan" in x["fila"]["texto"])
        IM.resolver(self.a, q["id"], {"accion": "ignorar"})
        self.assertEqual(len(self.a.todos("pendiente")), 0)
        # Reimportar el mismo extracto: nada nuevo, nada por revisar (lo resuelto y lo descartado se reconoce)
        r = self.importar(self.extracto())
        self.assertEqual((r["nuevas"], r["dudas"]), (0, 0))
        # Un extracto posterior con la misma ferretería se clasifica solo
        r = self.importar(self.extracto(FILAS + [("2026-09-20", "Pago Movil En Ferreteria Lopez, Madrid", -8.0)], nombre="b.xlsx"))
        self.assertEqual((r["nuevas"], r["dudas"]), (1, 0))
        self.assertEqual([m["categoria"] for m in self.movs() if m["fecha"] == "2026-09-20"], ["Hogar"])

    def test_recordar_resuelve_las_iguales(self):
        filas = [("2026-09-01", "Pago Movil En Bar Pepe, Madrid ES, Tarj. :*1234", -5.0), ("2026-09-02", "Pago Movil En Bar Pepe, Madrid ES, Tarj. :*1234", -7.0),
                 ("2026-09-03", "Compra Bet365, Palma, Tarjeta 4000000000001234 , Comision 0,00", -20.0)]
        self.a.guardar("regla", {"patron": "bet365", "categoria": "Ocio", "clase": "gasto"})
        r = self.importar(self.extracto(filas))
        self.assertEqual(r["dudas"], 0 if any(x["patron"] == "bar" for x in self.a.todos("regla")) else 2)
        self.assertEqual([m["categoria"] for m in self.movs() if "Bet365" in m["ext_texto"]], ["Ocio"])  # no «Comisiones» por «COMISION 0,00»
    def test_recordar_resuelve_las_iguales_pendientes(self):
        filas = [("2026-09-01", "Pago Movil En Ferreteria Lopez, Madrid", -5.0), ("2026-09-02", "Pago Movil En Ferreteria Lopez, Madrid", -7.0)]
        self.importar(self.extracto(filas))
        p = self.a.todos("pendiente")[0]
        msg = IM.resolver(self.a, p["id"], {"accion": "guardar", "clase": "gasto", "categoria": "Hogar", "recordar": True, "patron": "ferreteria lopez"})
        self.assertIn("1 más", msg)
        self.assertEqual((len(self.a.todos("pendiente")), sorted(m["importe"] for m in self.movs())), (0, [5.0, 7.0]))

    def test_traspaso_desconocido_a_otra_cuenta(self):
        self.importar(self.extracto([("2026-09-01", "Traspaso interno cuenta 9999", -100.0)]))
        p = self.a.todos("pendiente")[0]
        IM.resolver(self.a, p["id"], {"accion": "guardar", "clase": "transferencia", "cuenta_otra": "Bróker", "recordar": True, "patron": "cuenta 9999"})
        m = self.movs()[0]
        self.assertEqual((m["clase"], m.get("destino")), ("transferencia", "Bróker"))
        self.importar(self.extracto([("2026-09-01", "Traspaso interno cuenta 9999", -100.0), ("2026-10-01", "Traspaso interno cuenta 9999", -100.0)], nombre="c.xlsx"))
        self.assertEqual(sum(1 for m in self.movs() if m.get("destino") == "Bróker"), 2)

    def test_saldos_rotos_no_importa_nada(self):
        r = self.importar(self.extracto(romper=True))
        self.assertFalse(r["ok"])
        self.assertIn("no cuadran", r["mensaje"])
        self.assertEqual(len(self.movs()), 0)

    def test_saldos_con_los_dias_al_reves_dentro_de_cada_dia_se_ordenan(self):
        # Días de más reciente a más antiguo, pero cada día en orden cronológico: los saldos cuadran si se reordena dentro del día
        ruta = os.path.join(self.c.banco, "mixto.csv")
        escribir(ruta, "utf-8", "Fecha;Concepto;Importe;Saldo\n02/09/2026;Nomina Empresa SL;1500,00;2470,00\n02/09/2026;Bar Pepe;-30,00;2440,00\n"
                                "01/09/2026;Mercadona;-20,00;980,00\n01/09/2026;Lidl;-10,00;970,00\n")
        IM.crear_perfil(self.a, "Banco M", "banco", {"fecha": "Fecha", "concepto": "Concepto", "importe": "Importe", "saldo": "Saldo"}, "Nómina")
        r = self.importar(ruta)
        self.assertTrue(r["ok"], r.get("mensaje"))
        self.assertEqual(r["nuevas"] + r["dudas"], 4)
        self.assertEqual(self.a.config("saldo_extracto:Nómina"), {"fecha": "2026-09-02", "saldo": 2440.0})
        # Si dentro del día no hay orden posible, sigue sin importarse
        filas = [{"op": "2026-09-02", "texto": "A", "importe": 10.0, "saldo": 1010.0}, {"op": "2026-09-01", "texto": "B", "importe": -5.0, "saldo": 990.0}]
        self.assertIn("no cuadran", IM.ordenar_y_comprobar_saldos(filas)[1])

    def test_apuntado_a_mano_no_se_duplica(self):
        self.a.guardar("movimiento", {"fecha": "2026-09-01", "clase": "gasto", "categoria": "Supermercado", "importe": 45.2, "concepto": "Súper", "cuenta": "Nómina"})
        r = self.importar(self.extracto())
        self.assertEqual(r["nuevas"], 5)

    def test_apuntado_a_mano_dias_antes_se_enlaza(self):
        self.a.guardar("movimiento", {"fecha": "2026-09-03", "clase": "gasto", "categoria": "Hogar", "importe": 12.5, "concepto": "Tornillos", "cuenta": "Nómina"})
        self.a.guardar("movimiento", {"fecha": "2026-09-10", "clase": "ingreso", "categoria": "Otros ingresos", "importe": 10.0, "concepto": "No es la compra", "cuenta": "Nómina"})
        self.importar(self.extracto())
        tornillos = [m for m in self.movs() if m["importe"] == 12.5]
        self.assertEqual(len(tornillos), 1)
        self.assertEqual((tornillos[0]["concepto"], tornillos[0]["ext_fecha"]), ("Tornillos", "2026-09-05"))
        self.assertEqual(sum(1 for m in self.movs() if m["importe"] == 10.0), 3)  # el ingreso a mano no se confunde con las compras de 10 €
        r = self.importar(self.extracto(nombre="otra.xlsx"))
        self.assertEqual(r["nuevas"], 0)

    def test_formato_nuevo_csv(self):
        ruta = os.path.join(self.c.banco, "banco_x.csv")
        escribir(ruta, "utf-8", "Date;Description;Amount;Balance\n02/09/2026;CARREFOUR EXPRESS;-12,40;987,60\n01/09/2026;Recibo Iberdrola;-50,00;1.000,00\n")
        with self.assertRaises(IM.NecesitaPerfil) as e: self.importar(ruta)
        self.assertEqual(e.exception.info["cabecera"], ["Date", "Description", "Amount", "Balance"])
        IM.crear_perfil(self.a, "Banco X", "banco", {"fecha": "Date", "concepto": "Description", "importe": "Amount", "saldo": "Balance"}, "Nómina")
        r = self.importar(ruta)
        self.assertEqual((r["ok"], r["nuevas"]), (True, 2))
        self.assertEqual(sorted(m["categoria"] for m in self.movs()), ["Suministros", "Supermercado"])

    def test_cargo_y_abono(self):
        ruta = os.path.join(self.c.banco, "b.csv")
        escribir(ruta, "cp1252", "Fecha;Concepto;Cargo;Abono\n01/09/2026;Mercadona;23,10;\n03/09/2026;Nómina;;1500,00\n")
        IM.crear_perfil(self.a, "Banco Y", "banco", {"fecha": "Fecha", "concepto": "Concepto", "cargo": "Cargo", "abono": "Abono"}, "Nómina")
        r = self.importar(ruta)
        self.assertEqual(r["nuevas"], 2)
        self.assertEqual(sorted((m["clase"], m["importe"]) for m in self.movs()), [("gasto", 23.1), ("ingreso", 1500.0)])

    def test_las_filas_que_no_se_leen_se_avisan_y_lo_demas_se_importa(self):
        ruta = os.path.join(self.c.banco, "a_mano.csv")
        escribir(ruta, "utf-8", "Fecha;Concepto;Importe\n01/09/2026;Mercadona;-23,45\n2026-09-03T10:15:00Z;Lidl;-9,90\n05/09/2026;Kiosko;-3,2O\n31/02/2026;Farmacia;-7,00\n\nTotal;;-43,55\n")
        IM.crear_perfil(self.a, "Banco Z", "banco", {"fecha": "Fecha", "concepto": "Concepto", "importe": "Importe"}, "Nómina")
        r = self.importar(ruta)
        self.assertEqual((r["ok"], r["nuevas"] + r["dudas"]), (True, 2))  # la fecha con «T» se entiende: son dos movimientos
        # Las dos que no: un importe con una O por un cero y un día que no existe. La línea en blanco y el total no son movimientos.
        self.assertEqual([(x["fila"], x["que"], x["valor"], x["texto"]) for x in r["ilegibles"]], [(4, "el importe", "-3,2O", "Kiosko"), (5, "la fecha", "31/02/2026", "Farmacia")])
        self.assertIn("⚠ 2 filas no se han podido leer (la primera, la fila 4: no entiendo el importe «-3,2O»)", r["mensaje"])

    def test_con_saldo_una_fila_ilegible_no_cuadra_y_el_error_dice_cual(self):
        ruta = os.path.join(self.c.banco, "con_saldo.csv")
        escribir(ruta, "utf-8", "Fecha;Concepto;Importe;Saldo\n03/09/2026;Lidl;-10,00;956,55\n02/09/2026;Kiosko;-1O,00;966,55\n01/09/2026;Mercadona;-23,45;976,55\n")
        IM.crear_perfil(self.a, "Banco S", "banco", {"fecha": "Fecha", "concepto": "Concepto", "importe": "Importe", "saldo": "Saldo"}, "Nómina")
        r = self.importar(ruta)
        self.assertFalse(r["ok"])
        self.assertIn("no cuadran", r["mensaje"]); self.assertIn("en la fila 3 no entiendo el importe «-1O,00»", r["mensaje"])
        self.assertEqual(self.movs(), [])
        # y si no se lee ninguna, se dice por qué en vez de «el archivo no tiene movimientos»
        escribir(ruta, "utf-8", "Fecha;Concepto;Importe;Saldo\n2026/9;Lidl;-10,00;956,55\n")
        with self.assertRaises(ValueError) as e: self.importar(ruta)
        self.assertIn("no entiendo la fecha «2026/9» (fila 2)", str(e.exception))

    def test_manda_lo_que_el_archivo_es_por_dentro_no_su_extension(self):
        import openpyxl
        IM.crear_perfil(self.a, "Banco X", "banco", {"fecha": "Fecha", "concepto": "Concepto", "importe": "Importe"}, "Nómina")
        en = lambda nombre: os.path.join(self.c.banco, nombre)
        # un .xlsx al que el banco llama .xls
        wb = openpyxl.Workbook(); ws = wb.active
        ws.append(["Fecha", "Concepto", "Importe"]); ws.append([datetime.datetime(2026, 9, 1), "Mercadona", -23.45])
        wb.save(en("a.xlsx")); os.replace(en("a.xlsx"), en("a.xls"))
        # el «Texto Unicode» de Excel (UTF-16 con tabuladores) y un texto con tabuladores llamado .xls
        escribir(en("b.csv"), "utf-16", "Fecha\tConcepto\tImporte\n02/09/2026\tMercadona\t-11,10\n")
        escribir(en("c.xls"), "cp1252", "Fecha\tConcepto\tImporte\n03/09/2026\tMercadona\t-12,20\n")
        for nombre in ("a.xls", "b.csv", "c.xls"):
            r = self.importar(en(nombre))
            self.assertEqual((r["ok"], r["nuevas"]), (True, 1), nombre)
        self.assertEqual(sorted(m["importe"] for m in self.movs()), [11.1, 12.2, 23.45])
        # lo que no se puede leer se dice claro, no con el error de la librería
        escribir(en("d.xls"), "utf-8", "<html><body><table><tr><td>Fecha</td><td>Concepto</td><td>Importe</td></tr></table></body></html>")
        with io.open(en("e.xlsx"), "wb") as fh: fh.write(b"")
        with io.open(en("f.xls"), "wb") as fh: fh.write(b"\xd0\xcf\x11\xe0" + b"\x00" * 600)
        with io.open(en("g.xlsx"), "wb") as fh: fh.write(b"PK\x03\x04" + b"no es un excel")
        for nombre, texto in (("d.xls", "no es un Excel de verdad"), ("e.xlsx", "está vacío"), ("f.xls", "protegido con contraseña o dañado"), ("g.xlsx", "No se puede leer como Excel")):
            with self.assertRaises(ValueError) as e: self.importar(en(nombre))
            self.assertIn(texto, str(e.exception), nombre)

class TestAprender(Base):
    """Lo que la app aprende sola: tu historial, los grupos de «Por revisar» y los cambios de categoría."""
    def test_limpia_tarjetas_enmascaradas(self):
        from financebuddy import clasificar as C
        self.assertEqual(C.patron_sugerido("COMPRA TARJ. 5540XXXXXXXX1234 MERCADONA VALENCIA"), "mercadona valencia")
        self.assertEqual(C.patron_sugerido("RECIBO /VODAFONE ESPANA SAU"), "vodafone espana sau")
        self.assertEqual(C.comercio("TRANSACCION CONTACTLESS EN BAR PEPE 00123, MADRID ES"), "Bar Pepe")

    def test_recuerda_sin_regla(self):
        # La primera vez es duda; se resuelve SIN «recordar»; la siguiente vez ya no pregunta (lo saca del historial)
        self.importar(self.extracto([("2026-09-05", "Pago Movil En Ferreteria Lopez, Madrid", -12.5)]))
        p = self.a.todos("pendiente")[0]
        IM.resolver(self.a, p["id"], {"accion": "guardar", "clase": "gasto", "categoria": "Hogar"})
        r = self.importar(self.extracto([("2026-09-05", "Pago Movil En Ferreteria Lopez, Madrid", -12.5), ("2026-09-20", "Pago Movil En Ferreteria Lopez, Madrid", -8.0)], nombre="b.xlsx"))
        self.assertEqual((r["nuevas"], r["dudas"], r["aprendidas"]), (1, 0, 1))
        self.assertEqual([m["categoria"] for m in self.movs() if m["fecha"] == "2026-09-20"], ["Hogar"])
        self.assertFalse(any(x.get("origen") == "usuario" for x in self.a.todos("regla")))  # no ha hecho falta una regla

    def test_sugerencia_por_parecido(self):
        from financebuddy import clasificar as C
        mem = C.memoria([{"clase": "gasto", "categoria": "Hogar", "ext_texto": "Pago Movil En Ferreteria Lopez, Madrid"}])
        s = C.sugerir("Pago Movil En Ferreteria Lopes, Getafe", -9.0, mem)
        self.assertEqual((s["categoria"], s["clase"]), ("Hogar", "gasto"))
        self.assertIn("parecido", s["motivo"])
        self.assertIsNone(C.sugerir("Pago Movil En Joyeria Sol", -9.0, mem))
        s = C.sugerir("Pago Movil En Ferreteria Lopez, Madrid", 9.0, mem)  # te lo devuelven
        self.assertEqual(s["clase"], "reembolso")

    def test_resolver_un_grupo(self):
        filas = [("2026-09-01", "Pago Movil En Kiosko Ana, Madrid", -2.0), ("2026-09-02", "Pago Movil En Kiosko Ana, Madrid", -3.0),
                 ("2026-09-03", "Pago Movil En Joyeria Sol, Madrid", -30.0)]
        self.importar(self.extracto(filas))
        P = self.a.todos("pendiente")
        kiosko = [p["id"] for p in P if "Kiosko" in p["fila"]["texto"]]
        msg = IM.resolver(self.a, kiosko[0], {"accion": "guardar", "clase": "gasto", "categoria": "Ocio", "ids": kiosko})
        self.assertIn("1 más", msg)
        self.assertEqual([p["fila"]["texto"] for p in self.a.todos("pendiente")], ["Pago Movil En Joyeria Sol, Madrid"])

    def test_recategorizar_parecidos_y_recordar(self):
        filas = [("2026-09-01", "Compra Mercadona, Madrid", -10.0), ("2026-09-08", "Compra Mercadona, Madrid", -20.0), ("2026-09-09", "Compra Lidl, Madrid", -5.0)]
        self.importar(self.extracto(filas))
        m = next(x for x in self.movs() if x["importe"] == 10.0)
        self.assertEqual(IM.parecidos(self.a, m["id"])["n"], 1)
        msg = IM.recategorizar(self.a, m["id"], {"categoria": "Hogar", "parecidos": True, "recordar": True})
        self.assertIn("1 más", msg)
        self.assertEqual(sorted((x["ext_texto"][:14], x["categoria"]) for x in self.movs()),
                         [("Compra Lidl, M", "Supermercado"), ("Compra Mercado", "Hogar"), ("Compra Mercado", "Hogar")])
        regla = next(x for x in self.a.todos("regla") if x.get("origen") == "usuario")
        self.assertEqual((regla["patron"], regla["categoria"]), ("mercadona", "Hogar"))
        r = self.importar(self.extracto(filas + [("2026-09-20", "Compra Mercadona, Madrid", -7.0)], nombre="c.xlsx"))
        self.assertEqual(r["nuevas"], 1)
        self.assertEqual([x["categoria"] for x in self.movs() if x["fecha"] == "2026-09-20"], ["Hogar"])
        with self.assertRaises(ValueError): IM.recategorizar(self.a, m["id"], {"categoria": "No existe"})

    def test_parecidos_dice_si_ya_hay_una_regla_tuya_de_ese_comercio(self):
        """La pantalla marca sola «cambiar los demás y recordarlo» solo la primera vez: si ya hay una regla tuya de ese
        comercio, sale desactivado (lo que manda es la regla) y para eso necesita saber que existe y a dónde manda."""
        self.importar(self.extracto([("2026-09-01", "Compra Mercadona, Madrid", -10.0), ("2026-09-08", "Compra Mercadona, Madrid", -20.0)]))
        m = next(x for x in self.movs() if x["importe"] == 10.0)
        r = IM.parecidos(self.a, m["id"])
        self.assertEqual((r["n"], r["regla"]), (1, None))  # la primera vez no hay nada tuyo
        IM.recategorizar(self.a, m["id"], {"categoria": "Hogar", "parecidos": True, "recordar": True})
        r = IM.parecidos(self.a, m["id"])
        self.assertEqual(r["regla"]["categoria"], "Hogar")
        self.assertTrue(self.a.obtener("regla", r["regla"]["id"]))
        # Una regla de la plantilla no cuenta: sobre esa sí se puede poner la tuya desde ahí
        self.a.borrar("regla", r["regla"]["id"])
        self.assertIsNone(IM.parecidos(self.a, m["id"])["regla"])

    def test_categoria_icono_y_color(self):
        i = self.a.guardar("categoria", {"nombre": "Pádel", "grupo": "variable", "icono": "🎾", "color": "#12AB34"})
        self.assertEqual(self.a.obtener("categoria", i)["color"], "#12AB34")
        with self.assertRaises(ValueError): self.a.guardar("categoria", {"nombre": "X", "color": "red;}"})

class TestExtractosReales(Base):
    """Lo aprendido con extractos reales: titular e IBAN de la cabecera, traspasos cruzados, reglas nuevas, bróker."""
    def test_gasolina_no_es_suministros_y_bizum_no_reparte_recibos(self):
        filas = [("2026-08-20", "Bizum a favor de Ana Ruiz concepto gasofa", -10.0),
                 ("2026-09-02", "Compra Anthropic* Claude Sub, San Francisco, Tarjeta 5163830304139456 , Comision 0,00", -200.0),
                 ("2026-09-02", "Bizum de Pedro Gil concepto sin concepto", 7.0)]
        self.importar(self.extracto(filas))
        cat = {m["ext_texto"][:10]: (m["clase"], m.get("categoria")) for m in self.movs()}
        self.assertEqual(cat["Bizum a fa"], ("gasto", "Coche"))
        self.assertNotIn("Bizum de P", cat)  # sin gasto que repartir cerca (la suscripción no cuenta): se pregunta
        self.assertEqual(len(self.a.todos("pendiente")), 1)

    def test_reglas_nuevas_y_prefijos_de_pago(self):
        from financebuddy import clasificar as C
        reglas = C.ordenar_reglas(self.a.todos("regla"))
        cat = lambda t: (C.regla_para(t, reglas) or {}).get("categoria")
        self.assertEqual(cat("COMPRA BET365, PALMA DEMALLO, TARJETA 5163830304139456 , COMISION 0,00"), "Apuestas")
        self.assertEqual(cat("TRANSACCION CONTACTLESS EN E S EUROPA, ANDUJAR ES, TARJ. :*139456"), "Coche")
        self.assertEqual(cat("PAGO MOVIL EN BURGUER TOMAS, JAEN ES"), "Comer fuera")
        self.assertEqual(C.patron_sugerido("PAGO MOVIL EN SQ *ESTACION JAE, JAEN ES"), "estacion jae")
        self.assertEqual(C.patron_sugerido("COMPRA WL *STEAM PURCHASE, BELLEVUE"), "steam purchase")
        self.assertEqual(C.patron_sugerido("TRANSACCION CONTACTLESS EN 12229 PULL AND, JAEN ES"), "pull and")

    def test_plantilla_nueva_llega_a_instalaciones_existentes(self):
        cat = next(c for c in self.a.todos("categoria") if c["nombre"] == "Apuestas")
        self.a.borrar("categoria", cat["id"])
        for r in self.a.todos("regla"):
            if r["patron"] == "bet365": self.a.borrar("regla", r["id"])
        self.a.guardar("regla", {"patron": "mercadona", "categoria": "Hogar", "clase": "gasto", "origen": "usuario"})
        self.a.set_config("plantilla_version", 1)
        plantilla.instalar(self.a)
        self.assertTrue(any(c["nombre"] == "Apuestas" for c in self.a.todos("categoria")))
        self.assertTrue(any(r["patron"] == "bet365" for r in self.a.todos("regla")))
        self.assertEqual(sum(r["patron"] == "mercadona" for r in self.a.todos("regla")), 2)  # la del usuario sigue, no se duplica la de serie
        n = len(self.a.todos("regla")); plantilla.instalar(self.a); self.assertEqual(len(self.a.todos("regla")), n)

    def test_cuenta_por_iban_y_titular(self):
        self.a.guardar("cuenta", {"nombre": "Otra", "tipo": "corriente", "extracto": True})
        r = self.importar(self.extracto(FILAS, iban="ES12 0049 1111 2222 3333 4444", titular="GARCIA LOPEZ ANA"))
        self.assertTrue(r["ok"])
        self.assertEqual(next(c for c in self.a.todos("cuenta") if c["nombre"] == "Nómina")["iban"], "4444")
        self.assertEqual(self.a.config("titulares"), ["GARCIA LOPEZ ANA"])
        # El mismo formato con otro IBAN: no se da por hecho que sea la cuenta de antes
        with self.assertRaises(IM.NecesitaCuenta): IM.importar_archivo(self.a, self.c, self.extracto(FILAS, nombre="b.xlsx", iban="ES12 0049 1111 2222 3333 5555"), "banco")
        r = IM.importar_archivo(self.a, self.c, self.extracto(FILAS, nombre="b.xlsx", iban="ES12 0049 1111 2222 3333 5555"), "banco", "Otra")
        r = IM.importar_archivo(self.a, self.c, self.extracto(FILAS + [("2026-09-20", "Compra Lidl, Madrid", -3.0)], nombre="c.xlsx", iban="ES12 0049 1111 2222 3333 4444"), "banco")
        self.assertEqual((r["cuenta"], r["nuevas"]), ("Nómina", 1))  # sin preguntar: por el IBAN

    def test_traspaso_a_tu_nombre_se_cruza_con_el_broker(self):
        self.a.set_config("titulares", ["GARCIA LOPEZ ANA"])
        filas = [("2026-08-26", "Transferencia inmediata a favor de Ana García López concepto ahorro", -500.0),
                 ("2026-08-27", "Transferencia de GARCIA LOPEZ ANA, concepto gastos.", 300.0)]
        r = self.importar(self.extracto(filas))
        P = self.a.todos("pendiente")
        self.assertEqual({p["fila"]["clase"] for p in P}, {"transferencia"})  # a tu nombre → traspaso, no gasto/ingreso
        csv = os.path.join(self.c.inversion, "mi.csv")
        escribir(csv, "utf-8", "Fecha de operación;Fecha valor;Concepto;Importe\n26/08/2026;26/08/2026;ahorro;500,00\n")
        r = IM.importar_archivo(self.a, self.c, csv, "inversion", "Bróker")
        self.assertEqual((r["traspasos"], r["dudas"]), (1, 0))
        m = next(m for m in self.movs() if m["importe"] == 500.0)
        self.assertEqual((m["clase"], m.get("destino")), ("transferencia", "Bróker"))
        self.assertEqual([p["fila"]["importe"] for p in self.a.todos("pendiente")], [300.0])  # la otra sigue esperando su cuenta

    def test_broker_crea_activo_para_todo_el_grupo(self):
        from financebuddy import servidor
        csv = os.path.join(self.c.inversion, "mi.csv")
        escribir(csv, "utf-8", "Fecha de operación;Fecha valor;Concepto;Importe\n"
                 "04/09/2026;08/09/2026;FIDELITY PHYSICAL BITCOIN ET @;-55,43\n04/08/2026;06/08/2026;FIDELITY PHYSICAL BITCOIN ET @;-55,45\n"
                 "04/08/2026;06/08/2026;ETF ETFS Copper ETC @ 1;-50,31\n")
        IM.importar_archivo(self.a, self.c, csv, "inversion", "Bróker")
        P = self.a.todos("pendiente")
        sug = IM.sugerencia_inversion(next(p for p in P if "BITCOIN" in p["fila"]["texto"]), self.a.todos("activo"))
        self.assertEqual((sug["nuevo"], sug["clase"]), ("Fidelity Physical Bitcoin", "cripto"))
        cobre = next(p for p in P if "Copper" in p["fila"]["texto"])
        self.assertEqual(IM.sugerencia_inversion(cobre, [])["nuevo"], "ETFS Copper")
        ids = [p["id"] for p in P if "BITCOIN" in p["fila"]["texto"]]
        IM.resolver(self.a, ids[0], {"accion": "activo", "nuevo_activo": "Fidelity Physical Bitcoin", "ids": ids, "recordar": True, "patron": "fidelity physical bitcoin"})
        a = next(x for x in self.a.todos("activo") if x["nombre"] == "Fidelity Physical Bitcoin")
        self.assertEqual((a["clase"], a["patrones"]), ("cripto", ["fidelity physical bitcoin"]))
        self.assertEqual(sorted(x["importe"] for x in self.a.todos("aportacion")), [55.43, 55.45])

class TestSugeridos(Base):
    """Lo dudoso con categoría sugerida fiable se guarda ya (marcado «sugerido»); lo demás sigue en «Por revisar»."""
    def historial(self):
        self.a.guardar("movimiento", {"fecha": "2026-08-01", "clase": "gasto", "categoria": "Hogar", "importe": 9.0, "cuenta": "Nómina", "concepto": "Ferretería",
                                      "ext_texto": "Pago Movil En Ferreteria Lopez, Madrid", "ext_importe": -9.0, "ext_fecha": "2026-08-01"})
    FILAS_S = [("2026-09-05", "Pago Movil En Ferreteria Lopes, Madrid", -12.0),  # nombre parecido al del historial
               ("2026-09-06", "Pago Movil En Taller Raro Sur, Madrid", -30.0),  # sin pistas
               ("2026-09-07", "Traspaso interno cuenta 9999", -100.0)]  # un traspaso se sigue preguntando

    def test_se_guarda_ya_con_la_categoria_sugerida(self):
        self.historial()
        r = self.importar(self.extracto(self.FILAS_S))
        self.assertEqual((r["nuevas"], r["dudas"], r["sugeridas"]), (1, 2, 1))
        m = next(m for m in self.movs() if "Lopes" in m.get("ext_texto", ""))
        self.assertEqual((m["clase"], m["categoria"]), ("gasto", "Hogar"))
        self.assertIn("parecido", m["sugerido"])
        self.assertEqual(len(self.a.todos("pendiente")), 2)
        self.assertIn("1 con categoría sugerida", r["mensaje"])

    def test_apagado_lo_deja_por_revisar(self):
        self.historial(); self.a.set_config("guardar_sugeridos", False)
        r = self.importar(self.extracto(self.FILAS_S))
        self.assertEqual((r["nuevas"], r["dudas"], r["sugeridas"]), (0, 3, 0))
        self.assertFalse(any(m.get("sugerido") for m in self.movs()))

    def test_umbral_mas_alto_exige_mas(self):
        self.historial(); self.a.set_config("sugeridos_umbral", 1.0)
        self.assertEqual(self.importar(self.extracto(self.FILAS_S))["sugeridas"], 0)

    def test_lo_sugerido_no_ensena_ni_se_duplica(self):
        from financebuddy import clasificar as C
        self.historial(); self.importar(self.extracto(self.FILAS_S))
        mem = C.memoria(self.movs())
        self.assertNotIn("ferreteria lopes", mem)  # una suposición sin confirmar no aprende
        r = self.importar(self.extracto(self.FILAS_S))
        self.assertEqual((r["nuevas"], r["dudas"], r["existentes"]), (0, 0, 3))
        self.assertEqual(sum(1 for m in self.movs() if "Lopes" in m.get("ext_texto", "")), 1)

    def test_confirmar_y_cambiar(self):
        self.historial(); self.importar(self.extracto(self.FILAS_S))
        m = next(m for m in self.movs() if m.get("sugerido"))
        self.assertEqual(IM.confirmar_sugeridos(self.a, [m["id"]]), 1)
        self.assertFalse(self.a.obtener("movimiento", m["id"]).get("sugerido"))
        self.assertEqual(IM.confirmar_sugeridos(self.a), 0)
        # Cambiar la categoría también confirma
        self.a.set_config("sugeridos_umbral", 0.6)
        self.importar(self.extracto(self.FILAS_S + [("2026-09-08", "Pago Movil En Ferreteria Lopex, Madrid", -4.0)], nombre="b.xlsx"))
        n = next(m for m in self.movs() if "Lopex" in m.get("ext_texto", ""))
        self.assertTrue(n.get("sugerido"))
        IM.recategorizar(self.a, n["id"], {"categoria": "Compras"})
        self.assertEqual((self.a.obtener("movimiento", n["id"]).get("sugerido"), self.a.obtener("movimiento", n["id"])["categoria"]), (None, "Compras"))

    def test_vista_previa_cuenta_los_sugeridos_y_no_guarda(self):
        from financebuddy.servidor import App
        self.historial(); self.a.cerrar()
        app = App(self.dir)
        try:
            app.alm.guardar("cuenta", {"nombre": "Nómina", "tipo": "corriente", "extracto": True}) if not any(c["nombre"] == "Nómina" for c in app.alm.todos("cuenta")) else None
            ruta = self.extracto(self.FILAS_S)
            r = app._importar(ruta, "banco", "Nómina", previa=True)
            self.assertEqual(r["previa"]["sugeridos"], 1)
            self.assertTrue(any(x["sugerido"] for x in r["previa"]["muestra"]))
            self.assertEqual(sum(1 for m in app.alm.todos("movimiento") if m.get("sugerido")), 0)
            self.assertEqual(app.alm.todos("pendiente"), [])
        finally: app.alm.cerrar()
        self.a = Almacen(self.c.db)

class TestBizums(Base):
    """Bizums recibidos: la parte que te devuelven de un gasto que pagaste tú (bizums.py)."""
    def reparto(self, extra=()):
        # El 15 pagas 97,84 € en el súper; el 16 te mandan 4 Bizums de ~19,56 € (97,84 ÷ 5: tú incluido)
        filas = [("2026-09-10", "Pago Movil En Gimnasio Zeus, Jaen", -30.00),
                 ("2026-09-15", "Pago Movil En Mercadona Aguad, Jaen, Tarj. :*1234", -97.84),
                 ("2026-09-16", "Bizum de Ana Ruiz Lopez", 19.56), ("2026-09-16", "Bizum de Luis Perez Gil", 19.6),
                 ("2026-09-16", "Bizum de Eva Sanz Diaz", 19.56), ("2026-09-16", "Bizum de Pablo Cano Rey", 19.56)] + list(extra)
        return self.importar(self.extracto(sorted(filas)))
    def test_reparto_exacto_encuentra_el_gasto_aunque_no_sea_del_dia_anterior(self):
        self.reparto([("2026-09-12", "Pago Movil En Cine Plaza, Jaen", -40.00)])  # otro gasto en medio, más reciente
        self.assertEqual([p for p in self.a.todos("pendiente") if "Bizum" in p["fila"]["texto"]], [])
        reem = [m for m in self.movs() if m["clase"] == "reembolso"]
        self.assertEqual((len(reem), {m["categoria"] for m in reem}), (4, {"Supermercado"}))
        # Y quedan enlazados con el gasto: tu parte real = 97,84 − 78,28 = 19,56
        from financebuddy import bizums
        bizums.enlazar(self.a)
        super_ = next(m for m in self.movs() if m["clase"] == "gasto" and m["importe"] == 97.84)
        self.assertEqual({m["reembolsa"] for m in self.movs() if m["clase"] == "reembolso"}, {super_["id"]})
    def test_no_devuelve_mas_de_lo_que_costo(self):
        from financebuddy import bizums
        self.reparto()
        for i, quien in enumerate(("Mara Gil Ron", "Rosa Gil Ron")):
            self.a.guardar("movimiento", {"fecha": f"2026-09-1{7 + i}", "clase": "reembolso", "categoria": "Supermercado", "importe": 19.56, "concepto": f"Otro{i}",
                                          "cuenta": "Nómina", "ext_texto": f"Bizum de {quien}"})
        bizums.enlazar(self.a)
        extra = {m["concepto"]: m for m in self.movs() if m["concepto"].startswith("Otro")}
        self.assertIsNotNone(extra["Otro0"].get("reembolsa"))  # 5 × 19,56 = 97,80: cabe
        self.assertIsNone(extra["Otro1"].get("reembolsa"))     # el sexto ya no: no se devuelve más de lo que costó
    def test_los_hermanos_del_reparto_heredan_la_categoria(self):
        filas = [("2026-09-26", "Bizum de Ana Ruiz Lopez concepto comida", 22.12), ("2026-09-26", "Bizum de Luis Perez Gil", 22.12),
                 ("2026-09-26", "Bizum de Eva Sanz Diaz", 22.20), ("2026-09-27", "Bizum de Pablo Cano Rey", 22.12),
                 ("2026-09-26", "Bizum de Otro Distinto Mas", 5.00)]  # importe distinto: no es del reparto
        r = self.importar(self.extracto(sorted(filas)))
        self.assertEqual(r["dudas"], 1, r)
        self.assertEqual({m["categoria"] for m in self.movs() if m["importe"] > 20}, {"Comer fuera"})
        self.assertEqual(self.a.todos("pendiente")[0]["fila"]["importe"], 5.0)
    def test_gasto_de_otro_extracto(self):
        # El gasto está en un extracto anterior (ya importado) y los Bizums llegan en otro
        self.importar(self.extracto([("2026-09-10", "Pago Movil En Restaurante El Pino, Jaen", -60.00)], nombre="a.xlsx"))
        r = self.importar(self.extracto([("2026-09-13", "Bizum de Ana Ruiz Lopez", 15.00), ("2026-09-13", "Bizum de Luis Perez Gil", 15.00),
                                         ("2026-09-13", "Bizum de Eva Sanz Diaz", 15.00)], nombre="b.xlsx", saldo_inicial=940.0))
        self.assertEqual((r["dudas"], r["nuevas"]), (0, 3), r)  # 60 ÷ 4 = 15 (tú + 3)
        self.assertEqual({m["categoria"] for m in self.movs() if m["clase"] == "reembolso"}, {"Comer fuera"})
    def test_sin_gasto_sigue_preguntando(self):
        r = self.importar(self.extracto([("2026-09-13", "Bizum de Ana Ruiz Lopez", 15.00), ("2026-09-13", "Bizum de Luis Perez Gil", 15.00)]))
        self.assertEqual(r["dudas"], 2)
        self.assertIn("parece un reparto", self.a.todos("pendiente")[0]["duda"])
    def test_aritmetica(self):
        from financebuddy import bizums as B
        self.assertTrue(B.parte(97.84, 19.56, 5) and B.parte(97.84, 19.6, 5) and not B.parte(97.84, 20.5, 5))
        self.assertEqual(B.iguales(19.56, "2026-09-16", [{"fecha": "2026-09-16", "importe": 19.6}, {"fecha": "2026-09-17", "importe": 19.56}, {"fecha": "2026-09-20", "importe": 19.56}]), 2)

class TestCobros(Base):
    """Dividendos y comisiones de un activo (registro `cobro`)."""
    CSV = ("Fecha de operación;Fecha valor;Concepto;Importe\n"
           "10/09/2026;10/09/2026;FIDELITY MSCI WORLD INDEX;-150,00\n"
           "20/09/2026;20/09/2026;DIVIDENDO FIDELITY MSCI WORLD INDEX;3,40\n"
           "25/09/2026;25/09/2026;CUPON ISHARES GOLD ETC;1,10\n")
    def importar_csv(self):
        ruta = os.path.join(self.c.inversion, "d.csv"); escribir(ruta, "utf-8", self.CSV)
        return IM.importar_archivo(self.a, self.c, ruta, "inversion", "Bróker")
    def test_dividendo_de_un_activo_conocido_no_es_una_venta(self):
        r = self.importar_csv()  # (el activo «Fondo MSCI» ya existe en la base de las pruebas, con el patrón «msci world»)
        self.assertIn("1 dividendos", r["mensaje"], r)
        cobros = self.a.todos("cobro")
        self.assertEqual([(c["activo"], c["tipo"], c["importe"], c["fecha"]) for c in cobros], [("Fondo MSCI", "dividendo", 3.4, "2026-09-20")])
        self.assertEqual([x["importe"] for x in self.a.todos("aportacion")], [150.0])  # solo la compra: el dividendo no es venta
        # El cupón de un activo que aún no existe espera en «Por revisar», sugerido como dividendo
        p = next(p for p in self.a.todos("pendiente") if "CUPON" in p["fila"]["texto"])
        self.assertEqual(IM.sugerencia_inversion(p, self.a.todos("activo"))["accion"], "dividendo")
        # Reimportar no duplica
        ruta = os.path.join(self.c.inversion, "d2.csv"); escribir(ruta, "utf-8", self.CSV)
        IM.importar_archivo(self.a, self.c, ruta, "inversion", "Bróker")
        self.assertEqual(len(self.a.todos("cobro")), 1)
    def test_resolver_un_dividendo_y_arreglos_de_cartera(self):
        from financebuddy import cartera
        self.a.guardar("activo", {"nombre": "Oro", "clase": "etf", "cuenta": "Bróker"})
        self.importar_csv()
        p = next(p for p in self.a.todos("pendiente") if "CUPON" in p["fila"]["texto"])
        with self.assertRaises(ValueError): IM.resolver(self.a, p["id"], {"accion": "dividendo"})  # falta el activo
        IM.resolver(self.a, p["id"], {"accion": "dividendo", "activo": "Oro"})
        self.assertEqual(sorted((c["activo"], c["importe"]) for c in self.a.todos("cobro")), [("Fondo MSCI", 3.4), ("Oro", 1.1)])
        self.assertEqual(self.a.todos("pendiente"), [])
        # Renombrar propaga y unir lo pasa al destino; borrar se lleva los dividendos y no reaparecen
        oro = next(a for a in self.a.todos("activo") if a["nombre"] == "Oro")
        self.a.guardar("activo", {**oro, "nombre": "Oro físico"}, oro["id"])
        self.assertIn("Oro físico", [c["activo"] for c in self.a.todos("cobro")])
        ids = {a["nombre"]: a["id"] for a in self.a.todos("activo")}
        cartera.unir(self.a, ids["Oro físico"], ids["Fondo MSCI"])
        self.assertEqual({c["activo"] for c in self.a.todos("cobro")}, {"Fondo MSCI"})
        cartera.borrar(self.a, ids["Fondo MSCI"])
        self.assertEqual(self.a.todos("cobro"), [])
        self.assertEqual(len([i for i in self.a.todos("ignorado")]), 3)  # compra + 2 dividendos: no reaparecen

class TestParticipaciones(Base):
    def test_participaciones_y_tipo(self):
        self.assertEqual((IM.participaciones("ETF ETFS Copper ETC @ 2", -90.6), IM.participaciones("ETF ETFS Copper ETC @ 2", 90.6)), (-2.0, 2.0))
        self.assertIsNone(IM.participaciones("FIDELITY PHYSICAL BITCOIN ET @", 5))
        self.assertEqual([IM.clase_activo(t) for t in ("ISHARES PHYSICAL GOLD ETC", "ETF ETFS Copper ETC", "FIDELITY S&P 500 INDEX P ACC")], ["materia", "materia", "fondo"])
        self.a.guardar("activo", {"nombre": "Cobre", "clase": "materia", "cuenta": "Bróker", "patrones": ["copper"], "aportado_inicial": 0})
        csv = os.path.join(self.c.inversion, "c.csv")
        escribir(csv, "utf-8", "Fecha de operación;Fecha valor;Concepto;Importe\n27/04/2026;29/04/2026;ETF ETFS Copper ETC @ 2;-90,63\n26/05/2026;28/05/2026;ETF ETFS Copper ETC @ 1;47,10\n")
        IM.importar_archivo(self.a, self.c, csv, "inversion", "Bróker")
        self.assertEqual(sorted((x["importe"], x["participaciones"]) for x in self.a.todos("aportacion")), [(-47.1, -1.0), (90.63, 2.0)])

class TestInversion(Base):
    CSV = ("Fecha de operación;Fecha valor;Concepto;Importe\n"
           "01/09/2026;01/09/2026;Ahorro;200,00\n"
           "10/09/2026;10/09/2026;FIDELITY MSCI WORLD INDEX;-150,00\n"
           "12/09/2026;12/09/2026;ISHARES GOLD ETC;-50,00\n"
           "30/09/2026;30/09/2026;Periodo 01/09/2026 30/09/2026;0,45\n")
    def csv(self, nombre="mi.csv"):
        ruta = os.path.join(self.c.inversion, nombre); escribir(ruta, "utf-8", self.CSV); return ruta
    def test_importa_y_aprende(self):
        r = IM.importar_archivo(self.a, self.c, self.csv(), "inversion", "Bróker")
        self.assertEqual((r["nuevas"], r["dudas"]), (2, 2))  # compra MSCI + intereses; dudas: entrada y oro
        self.assertEqual([(x["activo"], x["importe"]) for x in self.a.todos("aportacion")], [("Fondo MSCI", 150.0)])
        self.assertEqual([(m["clase"], m["importe"], m["cuenta"]) for m in self.movs()], [("ingreso", 0.45, "Bróker")])
        self.a.guardar("activo", {"nombre": "Oro", "clase": "etf", "cuenta": "Bróker"})
        for p in self.a.todos("pendiente"):
            if "GOLD" in p["fila"]["texto"]: IM.resolver(self.a, p["id"], {"accion": "activo", "activo": "Oro", "recordar": True, "patron": "ishares gold"})
            else: IM.resolver(self.a, p["id"], {"accion": "ignorar", "recordar": True, "patron": "ahorro"})
        r = IM.importar_archivo(self.a, self.c, self.csv("mi2.csv"), "inversion", "Bróker")
        self.assertEqual((r["nuevas"], r["dudas"]), (0, 0))
        self.assertIn("ishares gold", next(a for a in self.a.todos("activo") if a["nombre"] == "Oro")["patrones"])

    def test_sin_broker_no_ofrece_otras_cuentas(self):
        self.a.borrar("cuenta", next(c for c in self.a.todos("cuenta") if c["tipo"] == "broker")["id"])
        with self.assertRaises(IM.NecesitaCuenta) as e: IM.importar_archivo(self.a, self.c, self.csv(), "inversion")
        self.assertEqual(e.exception.info["cuentas"], [])

class TestOperaciones(Base):
    """Órdenes de fondos (CSV de MyInvestor: ISIN, importe, participaciones, estado) y Excel de operaciones con títulos."""
    ORDENES = ("Fecha de la orden;ISIN;Importe estimado;Nº de participaciones;Estado\n"
               "12/03/2026;IE00BYX5MX67;500.2 EUR;30,5;Finalizada\n"
               "10/03/2026;FR0000989626;500 EUR;0,011;Finalizada\n"
               "02/04/2026;LU9999999991;80 EUR;1,5;Finalizada\n"
               "01/04/2026;LU0034353002;300 EUR;;Cancelada\n"
               "04/03/2026;FR0000989626;1000 EUR;0,022;Finalizada\n")
    CUENTA = "Fecha de operación;Fecha valor;Concepto;Importe\n06/03/2026;06/03/2026;GROUPAMA TRESORERIE I ACC EUR;-999,50\n"
    def archivo(self, nombre, texto, enc="cp1252"):
        ruta = os.path.join(self.c.inversion, nombre); escribir(ruta, enc, texto); return ruta
    def aps(self, activo=None): return sorted((a["fecha"], a["activo"], a["importe"], a.get("participaciones")) for a in self.a.todos("aportacion") if not activo or a["activo"] == activo)

    def test_ordenes_y_traspaso_entre_fondos(self):
        r = IM.importar_archivo(self.a, self.c, self.archivo("ordenes.csv", self.ORDENES), None, "Bróker")
        self.assertEqual(r["tipo"], "operaciones"); self.assertEqual(r["nuevas"], 4)
        self.assertIn("1 traspasos entre fondos", r["mensaje"]); self.assertIn("1 canceladas", r["mensaje"])
        self.assertEqual(self.aps(), [("2026-03-04", "Groupama Trésorerie", 1000.0, 0.022), ("2026-03-10", "Groupama Trésorerie", -500.0, -0.011),
                                      ("2026-03-12", "Fidelity S&P 500", 500.2, 30.5), ("2026-04-02", "Fondo LU9999999991", 80.0, 1.5)])
        A = {a["nombre"]: a for a in self.a.todos("activo")}
        self.assertEqual((A["Fidelity S&P 500"]["isin"], A["Fondo LU9999999991"]["clase"]), ("IE00BYX5MX67", "fondo"))
        tr = [a["activo"] for a in self.a.todos("aportacion") if a.get("traspaso")]
        self.assertEqual(sorted(tr), ["Fidelity S&P 500", "Groupama Trésorerie"])
        # Otra vez el mismo archivo: nada nuevo
        r = IM.importar_archivo(self.a, self.c, self.archivo("ordenes2.csv", self.ORDENES), None, "Bróker")
        self.assertEqual((r["nuevas"], r["existentes"]), (0, 4))
        # El extracto de la cuenta trae la compra de Groupama: es la misma, no se duplica (y deja de ser «supuesta»)
        r = IM.importar_archivo(self.a, self.c, self.archivo("cuenta.csv", self.CUENTA), "inversion", "Bróker")
        self.assertEqual((r["nuevas"], r["existentes"], r["dudas"]), (0, 1, 0))
        g = next(a for a in self.a.todos("aportacion") if a["fecha"] == "2026-03-04")
        self.assertEqual((g["ext_importe"], g.get("supuesta", "")), (-999.5, ""))

    def test_cuenta_antes_que_las_ordenes(self):
        r = IM.importar_archivo(self.a, self.c, self.archivo("cuenta.csv", self.CUENTA), "inversion", "Bróker")
        self.assertEqual(r["dudas"], 1)  # aún no hay activo que reconozca «Groupama»
        IM.importar_archivo(self.a, self.c, self.archivo("ordenes.csv", self.ORDENES), None, "Bróker")
        self.assertEqual(self.a.todos("pendiente"), [])  # la duda era esa orden
        self.assertEqual(self.aps("Groupama Trésorerie"), [("2026-03-06", "Groupama Trésorerie", 999.5, 0.022), ("2026-03-10", "Groupama Trésorerie", -500.0, -0.011)])

    def test_venta_a_la_cuenta_corrige_la_orden_supuesta(self):
        IM.importar_archivo(self.a, self.c, self.archivo("o.csv", "Fecha de la orden;ISIN;Importe estimado;Nº de participaciones;Estado\n"
                                                                 "14/04/2026;LU0034353002;199.98 EUR;2,199;Finalizada\n"), None, "Bróker")
        self.assertEqual(self.aps(), [("2026-04-14", "DWS Floating Rate Notes", 199.98, 2.199)])  # sin más datos, se supone compra
        IM.importar_archivo(self.a, self.c, self.archivo("c.csv", "Fecha de operación;Fecha valor;Concepto;Importe\n"
                                                                 "16/04/2026;17/04/2026;DEUTSCHE FLOAT RATE NOTS LC EU;199,98\n"), "inversion", "Bróker")
        self.assertEqual(self.aps(), [("2026-04-14", "DWS Floating Rate Notes", -199.98, -2.199)])  # el dinero entró en la cuenta: era una venta

    def test_excel_de_titulos(self):
        import openpyxl
        ruta = os.path.join(self.c.inversion, "operaciones.xlsx")
        wb = openpyxl.Workbook(); ws = wb.active
        ws.append(["Fecha", "Tipo", "Activo", "Estado", "Títulos", "Títulos netos"])
        ws.append([datetime.datetime(2026, 6, 24), "Compra", "FIDELITY PHYSICAL BITCOIN ET", "Finalizada", 57, "=E2"])
        ws.append([datetime.datetime(2026, 7, 30), "Compra", "FIDELITY PHYSICAL BITCOIN ET", "Rechazada", 3, "=E3"])
        ws.append([datetime.datetime(2026, 5, 25), "Venta", "ISHARES PHYSICAL GOLD ETC", "Finalizada", 1, "=E4"])
        wb.save(ruta)
        r = IM.importar_archivo(self.a, self.c, ruta, None, "Bróker")
        self.assertIn("2 esperando al extracto", r["mensaje"])
        self.assertEqual(sorted(a["nombre"] for a in self.a.todos("activo")), ["Fidelity Physical Bitcoin", "Fondo MSCI", "iShares Physical Gold"])
        r = IM.importar_archivo(self.a, self.c, self.archivo("c.csv", "Fecha de operación;Fecha valor;Concepto;Importe\n"
                                                                 "25/06/2026;29/06/2026;FIDELITY PHYSICAL BITCOIN ET @;-303,41\n"
                                                                 "26/05/2026;28/05/2026;ISHARES PHYSICAL GOLD ETC @ 1;74,59\n"), "inversion", "Bróker")
        self.assertIn("completadas con sus títulos", r["mensaje"])
        self.assertEqual(self.aps(), [("2026-05-26", "iShares Physical Gold", -74.59, -1.0), ("2026-06-25", "Fidelity Physical Bitcoin", 303.41, 57.0)])
        self.assertEqual(self.a.todos("operacion"), [])

    def test_una_orden_con_la_fecha_ilegible_se_avisa(self):
        ruta = self.archivo("ordenes.csv", "Fecha;Tipo;Activo;Estado;Títulos\n24/06/2026;Compra;FIDELITY PHYSICAL BITCOIN ET;Finalizada;57\n2026-13-45;Compra;ISHARES PHYSICAL GOLD ETC;Finalizada;2\n")
        r = IM.importar_archivo(self.a, self.c, ruta, None, "Bróker")
        self.assertEqual([(x["fila"], x["valor"]) for x in r["ilegibles"]], [(3, "2026-13-45")])
        self.assertIn("⚠ 1 fila no se ha podido leer (la fila 3: no entiendo la fecha «2026-13-45»)", r["mensaje"])

class TestCartera(Base):
    """Arreglos a mano de la cartera (cartera.py) y sugerencias del bróker que no inventan activos."""
    CUENTA = ("Fecha de operación;Fecha valor;Concepto;Importe\n"
              "04/03/2025;04/03/2025;Inicio;5000\n"
              "06/03/2025;06/03/2025;GROUPAMA TRESORERIE I ACC EUR;-4999,04\n"
              "18/03/2025;19/03/2025;S&P 500 INDEX P ACC EUR @ 31.5;-400\n")
    def importar_cuenta(self, nombre="c.csv"):
        ruta = os.path.join(self.c.inversion, nombre); escribir(ruta, "utf-8", self.CUENTA)
        return IM.importar_archivo(self.a, self.c, ruta, "inversion", "Bróker")
    def test_sugerencias_del_broker(self):
        self.importar_cuenta()
        s = {p["fila"]["texto"]: IM.sugerencia_inversion(p, self.a.todos("activo")) for p in self.a.todos("pendiente")}
        self.assertEqual(s["Inicio"], {"accion": "ignorar"})  # dinero desde tu banco, no la venta de un activo «Inicio»
        self.assertEqual(s["GROUPAMA TRESORERIE I ACC EUR"], {"accion": "activo", "nuevo": "Groupama Trésorerie", "clase": "fondo", "isin": "FR0000989626"})
        g = next(p for p in self.a.todos("pendiente") if p["fila"]["texto"].startswith("GROUPAMA"))
        IM.resolver(self.a, g["id"], {"accion": "activo", "nuevo_activo": "Groupama Trésorerie"})
        act = next(a for a in self.a.todos("activo") if a["nombre"] == "Groupama Trésorerie")
        self.assertEqual((act["isin"], act["clase"]), ("FR0000989626", "fondo"))  # así las órdenes (por ISIN) caen en el mismo

    def test_era_un_traspaso(self):
        from financebuddy import cartera
        self.importar_cuenta()
        p = next(p for p in self.a.todos("pendiente") if p["fila"]["texto"] == "Inicio")
        IM.resolver(self.a, p["id"], {"accion": "activo", "nuevo_activo": "Inicio"})  # el error de antes: tomarlo por una venta
        ini = next(a for a in self.a.todos("activo") if a["nombre"] == "Inicio")
        cartera.borrar(self.a, ini["id"], era_traspaso=True)
        self.assertFalse(any(a["nombre"] == "Inicio" for a in self.a.todos("activo")))
        self.assertFalse(any(x["activo"] == "Inicio" for x in self.a.todos("aportacion")))
        r = self.importar_cuenta("c2.csv")  # el mismo extracto otra vez: no vuelve
        self.assertFalse(any(p["fila"]["texto"] == "Inicio" for p in self.a.todos("pendiente")))

    def test_unir_y_cuadrar(self):
        from financebuddy import cartera
        self.a.guardar("activo", {"nombre": "Sp500", "clase": "otro", "cuenta": "Bróker", "valor": 100, "fecha_valor": "2025-03-20"})
        self.a.guardar("activo", {"nombre": "Fidelity S&P 500", "clase": "fondo", "cuenta": "Bróker", "patrones": ["fidelity s&p 500"]})
        self.a.guardar("aportacion", {"fecha": "2025-03-18", "activo": "Sp500", "importe": 400, "participaciones": 31.5, "cuenta": "Bróker"})
        self.a.guardar("aportacion", {"fecha": "2025-04-08", "activo": "Fidelity S&P 500", "importe": 198.7, "participaciones": 17.807, "cuenta": "Bróker"})
        self.a.guardar("patrimonio", {"fecha": "2025-03-31", "saldos": {"Bróker": 0}, "valores": {"Sp500": 101, "Fidelity S&P 500": 0}})
        A = {a["nombre"]: a["id"] for a in self.a.todos("activo")}
        cartera.unir(self.a, A["Sp500"], A["Fidelity S&P 500"])
        self.assertEqual(sorted(a["nombre"] for a in self.a.todos("activo")), ["Fidelity S&P 500", "Fondo MSCI"])
        f = next(a for a in self.a.todos("activo") if a["nombre"] == "Fidelity S&P 500")
        self.assertIn("sp500", f["patrones"]); self.assertEqual(f["valor"], 100)
        self.assertEqual(self.a.todos("patrimonio")[0]["valores"], {"Fidelity S&P 500": 101})
        self.assertAlmostEqual(cartera.participaciones(self.a, "Fidelity S&P 500"), 49.307)
        # El bróker dice 49,326: se añade un ajuste de +0,019 sin dinero; la segunda vez ya cuadra
        self.assertIn("+0.019", cartera.cuadrar(self.a, f["id"], "49,326", "2025-05-01", "650"))
        aj = [x for x in self.a.todos("aportacion") if x.get("ajuste")]
        self.assertEqual([(x["importe"], x["participaciones"]) for x in aj], [(0.0, 0.019)])
        self.assertIn("ya cuadraba", cartera.cuadrar(self.a, f["id"], 49.326))
        self.assertEqual(next(a for a in self.a.todos("activo") if a["nombre"] == "Fidelity S&P 500")["valor"], 650)

    def test_ordenes_mandan_sobre_participaciones_redondeadas(self):
        self.importar_cuenta()
        s = next(p for p in self.a.todos("pendiente") if p["fila"]["texto"].startswith("S&P"))
        IM.resolver(self.a, s["id"], {"accion": "activo", "nuevo_activo": "S&P 500 Index P Acc Eur"})
        ruta = os.path.join(self.c.inversion, "o.csv")
        escribir(ruta, "cp1252", "Fecha de la orden;ISIN;Importe estimado;Nº de participaciones;Estado\n16/03/2025;IE00BYX5MX67;400 EUR;31,519;Finalizada\n")
        IM.importar_archivo(self.a, self.c, ruta, None, "Bróker")
        self.assertEqual([(x["activo"], x["participaciones"]) for x in self.a.todos("aportacion")], [("S&P 500 Index P Acc Eur", 31.519)])

class TestDetectar(Base):
    MESES = ["2026-07", "2026-08", "2026-09"]
    def filas(self):
        out = []
        for i, m in enumerate(self.MESES):
            out += [(f"{m}-01", "Transferencia inmediata a favor de Pedro Casero concepto Alquiler", -600.0),
                    (f"{m}-03", "Compra Mercadona, Madrid", -40.0), (f"{m}-12", "Compra Mercadona, Madrid", -55.0 - i),
                    (f"{m}-05", "Bizum a favor de Carlos Lopez concepto: wifi", -16.0),
                    (f"{m}-{27 - i:02d}", "Transferencia de Empresa SL, concepto Nomina mes", 1800.0),
                    (f"{m}-{29 - i:02d}", "Recibo Digi Spain Telecom, concepto: factura", -20.0 - i * 0.5)]
        out.append(("2026-09-15", "Transferencia de Club Deportivo, concepto pago", 300.0))
        return sorted(out)
    def preparar(self):
        self.importar(self.extracto(self.filas(), saldo_inicial=3000))
        for p in self.a.todos("pendiente"):  # el alquiler por transferencia llega como duda
            IM.resolver(self.a, p["id"], {"accion": "guardar", "clase": p["fila"]["clase"] if p["fila"]["importe"] > 0 else "gasto",
                                          "categoria": "Vivienda" if "casero" in p["fila"]["texto"].lower() else "Otros ingresos" if p["fila"]["importe"] > 0 else "Otros"})
    def test_detecta_fijos_y_origenes(self):
        from financebuddy import detectar as D
        self.preparar()
        f = {x["patron"]: x for x in D.fijos(self.a)}
        self.assertEqual(set(f), {"empresa sl", "pedro casero", "carlos lopez concepto: wifi", "digi spain telecom"})
        self.assertEqual((f["empresa sl"]["categoria"], f["empresa sl"]["grupo"], f["empresa sl"]["importe"]), ("Nómina", "ingreso", 1800.0))
        self.assertEqual((f["digi spain telecom"]["categoria"], f["digi spain telecom"]["grupo"]), ("Suministros", "fijo"))
        self.assertEqual((f["pedro casero"]["dia"], f["pedro casero"]["desde"]), (1, "2026-09-01"))
        o = {x["origen"]: x for x in D.origenes(self.a)}
        self.assertEqual((o["Empresa SL"]["total"], o["Empresa SL"]["veces"], o["Empresa SL"]["media_mes"]), (5400.0, 3, 1800.0))
        self.assertIn("Club Deportivo", o)

    def test_crear_enlaza_y_no_repite(self):
        from financebuddy import detectar as D
        self.preparar()
        props = [x for x in D.fijos(self.a) if x["patron"] in ("empresa sl", "carlos lopez concepto: wifi")]
        self.assertIn("2 fijos creados", D.crear(self.a, props))
        recs = {r["nombre"]: r for r in self.a.todos("recurrente")}
        self.assertEqual(set(recs), {x["nombre"] for x in props})
        enlazados = [m for m in self.movs() if m.get("recurrente")]
        self.assertEqual(sorted(m["fecha"][:7] for m in enlazados), ["2026-09", "2026-09"])
        self.assertEqual({x["patron"] for x in D.fijos(self.a)}, {"pedro casero", "digi spain telecom"})
        with self.assertRaises(ValueError): D.crear(self.a, props[:1])
        # El mes siguiente se reconoce solo y queda enlazado
        self.importar(self.extracto(self.filas() + [("2026-10-27", "Transferencia de Empresa SL, concepto Nomina oct", 1850.0)], nombre="oct.xlsx", saldo_inicial=3000))
        oct_ = next(m for m in self.movs() if m["fecha"] == "2026-10-27")
        self.assertEqual((oct_["recurrente"], oct_["categoria"]), (props[0]["nombre"] if props[0]["patron"] == "empresa sl" else props[1]["nombre"], "Nómina"))

class TestBase(Base):
    def test_renombrar_propaga(self):
        self.importar(self.extracto())
        c = next(x for x in self.a.todos("cuenta") if x["nombre"] == "Ahorro")
        self.a.guardar("cuenta", {**c, "nombre": "Colchón"}, c["id"])
        self.assertTrue(any(m.get("destino") == "Colchón" for m in self.movs()))
    def test_instantanea_y_recuperar(self):
        foto = self.a.instantanea()
        self.importar(self.extracto())
        self.assertTrue(self.movs())
        self.a.recuperar(foto)
        self.assertFalse(self.movs())
    def test_nombres_unicos(self):
        with self.assertRaises(ValueError): self.a.guardar("cuenta", {"nombre": "nómina".capitalize().replace("o", "ó")})
    def test_validacion(self):
        with self.assertRaises(ValueError): self.a.guardar("movimiento", {"fecha": "ayer", "importe": 3, "concepto": "x"})
        with self.assertRaises(ValueError): self.a.guardar("movimiento", {"fecha": "2026-01-01", "importe": -3, "concepto": "x"})
    def test_copias(self):
        p = self.a.copia(self.c.copias, forzar=True)
        self.assertTrue(os.path.exists(p))

class TestLectura(unittest.TestCase):
    def test_numeros(self):
        for txt, n in [("1.234,56", 1234.56), ("-12,30", -12.3), ("1,234.56", 1234.56), ("1234.5", 1234.5), ("12,30 €", 12.3), ("(5,00)", -5.0), ("abc", None), ("", None), ("1.234.567", 1234567)]:
            self.assertEqual(L.numero(txt), n, txt)
    def test_fechas(self):
        for txt in ("05/09/2026", "2026-09-05", "05-09-2026", "05/09/26", datetime.date(2026, 9, 5)):
            self.assertEqual(L.fecha(txt), "2026-09-05")
    def test_fechas_con_hora_con_zona_o_con_el_mes_escrito(self):
        for txt in ("2026-09-05T10:15:23Z", "2026-09-05T10:15:23.160+02:00", "2026-09-05 10:15", "05/09/2026, 10:15", "05.09.26", "5 sep 2026", "5 de Septiembre de 2026",
                    "05-sept.-2026", "Sep 5, 2026", "September 5 2026"):
            self.assertEqual(L.fecha(txt), "2026-09-05", txt)
        # Lo ambiguo o imposible sigue sin ser una fecha: orden americano, todo junto, un día que no existe, un rótulo con cifras
        for txt in ("09/30/2026", "20260905", "31/02/2026", "5 martes 2026", "30/09/2026 | 10:00:00", "Saldo a 30/09/2026", "12,50", "1.234,56"):
            self.assertIsNone(L.fecha(txt), txt)
    def test_el_signo_menos_tipografico(self):
        self.assertEqual((L.numero("−12,30"), L.numero("–1.234,56 €"), L.numero("–")), (-12.3, -1234.56, None))
    def test_punto_de_miles_sin_decimales(self):
        # «1.234» es mil doscientos treinta y cuatro (punto de miles español); lo demás se lee como antes
        for txt, n in [("1.234", 1234), ("-1.500", -1500), ("1.500-", -1500), ("1.000 €", 1000), ("(2.000)", -2000), ("999.999", 999999),
                       ("1.234,56", 1234.56), ("12.5", 12.5), ("12.50", 12.5), ("1,234.56", 1234.56), ("1234.567", 1234.567), ("0.123", 0.123),
                       ("1.2345", 1.2345), ("1,234", 1.234), ("1.234.567", 1234567)]:
            self.assertEqual(L.numero(txt), n, txt)
        self.assertEqual(L.numero("1.234", miles=False), 1.234)
    def test_punto_decimal_se_decide_por_archivo(self):
        self.assertTrue(L.punto_decimal(["1.234", "12.50"]))  # «12.50»: el punto es la coma decimal
        self.assertTrue(L.punto_decimal(["199.98 EUR"]))
        self.assertFalse(L.punto_decimal(["1.234", "-45,20", "1.234,56", 12.5, None, ""]))
    def test_numeros_de_excel_que_no_son_fechas(self):
        self.assertIsNone(L.fecha(25000))  # un importe entero no es una fecha de 1968
        self.assertIsNone(L.fecha(80000))
        self.assertEqual(L.fecha(46295), "2026-09-30")  # número de serie de Excel de verdad
        self.assertEqual(L.fecha(46295.5), "2026-09-30")
        # Al adivinar columnas, una de números enteros no se toma por fechas sin ninguna fecha de verdad al lado
        c = L.columnas_probables(["A", "B", "C"], [[46200, "Compra Lidl", 46100], [46201, "Nómina", 45000]])
        self.assertEqual(c, {})
        c = L.columnas_probables(["A", "B", "C"], [["01/09/2026", "Compra Lidl", 46100], ["02/09/2026", "Nómina", 45000]])
        self.assertEqual((c["fecha"], c["concepto"], c["importe"]), ("A", "B", "C"))
        # con el nombre «Fecha» en la cabecera, los números de serie sí valen
        c = L.columnas_probables(["Fecha", "Concepto", "Importe"], [[46200, "Compra Lidl", -12.5]])
        self.assertEqual(c["fecha"], "Fecha")

class TestMilesAlImportar(Base):
    PERFIL = {"nombre": "Banco de prueba", "tipo": "banco", "columnas": {"fecha": "fecha", "concepto": "concepto", "importe": "importe", "saldo": "saldo"}}
    def leer(self, texto):
        self.a.guardar("perfil", self.PERFIL)
        ruta = os.path.join(self.c.banco, "b.csv"); escribir(ruta, "utf-8", texto)
        return [(f["importe"], f.get("saldo")) for f in IM.leer(ruta, self.a, "banco", "Banco de prueba")[1]]

    def test_formato_espanol(self):
        self.assertEqual(self.leer("Fecha;Concepto;Importe;Saldo\n01/09/2026;Nómina;1.500;2.734,56\n02/09/2026;Lidl;-12,30;2.722,26\n"),
                         [(1500.0, 2734.56), (-12.3, 2722.26)])
    def test_archivo_con_punto_decimal_no_cambia(self):
        # si el archivo escribe «12.50», su «1.234» es 1,234 como siempre (no se adivina fila a fila)
        self.assertEqual(self.leer("Fecha;Concepto;Importe;Saldo\n01/09/2026;Café;-1.234;100.50\n02/09/2026;Lidl;-12.30;88.20\n"),
                         [(-1.23, 100.5), (-12.3, 88.2)])
    def test_participaciones_con_punto_son_decimales(self):
        ruta = os.path.join(self.c.inversion, "o.csv")
        escribir(ruta, "utf-8", "Fecha de la orden;ISIN;Importe estimado;Nº de participaciones;Estado\n"
                                "14/04/2026;LU0034353002;1.500 EUR;2.199;Finalizada\n")
        from financebuddy import operaciones as OP
        ops = OP.leer(ruta, self.a)[1]
        self.assertEqual([(o["importe"], o["part"]) for o in ops], [(1500.0, 2.199)])

if __name__ == "__main__":
    unittest.main()
