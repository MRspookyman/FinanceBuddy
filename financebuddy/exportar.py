# Lo que sale de la app hacia un archivo tuyo: la plantilla de Excel para apuntar o pasar movimientos con un desplegable de TUS
# categorías (el importador la reconoce: perfil «Plantilla de FinanceBuddy»). El resumen en HTML se monta en la página
# (web/paneles/exportar.js) porque ahí están ya todas las cuentas; aquí no hace falta tocar nada de eso.
import io
from . import VERSION

def plantilla_excel(alm):
    """Excel (bytes) con una hoja «Movimientos» (Fecha, Concepto, Importe, Categoría con desplegable), otra «Listas» con tus
    categorías y otra con las instrucciones. Importe: negativo = gasto, positivo = ingreso."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation
    cats = sorted((c for c in alm.todos("categoria") if not c.get("oculta")), key=lambda c: (c.get("grupo") == "ingreso", c["nombre"].lower()))
    wb = Workbook()
    ws = wb.active; ws.title = "Movimientos"
    ws.append(["Fecha", "Concepto", "Importe", "Categoría"])
    for c in ws[1]: c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="5E8266"); c.alignment = Alignment(vertical="center")
    for col, ancho in zip("ABCD", (14, 42, 14, 24)): ws.column_dimensions[col].width = ancho
    ws.freeze_panes = "A2"
    for r in range(2, 1002):
        ws.cell(r, 1).number_format = "dd/mm/yyyy"; ws.cell(r, 3).number_format = "#,##0.00"
    ls = wb.create_sheet("Listas")
    ls.append(["Categorías"]); ls["A1"].font = Font(bold=True)
    for c in cats: ls.append([c["nombre"]])
    ls.column_dimensions["A"].width = 26
    n = max(2, len(cats) + 1)
    dv = DataValidation(type="list", formula1=f"Listas!$A$2:$A${n}", allow_blank=True, showErrorMessage=True,
                        errorTitle="Categoría", error="Elige una de la lista (las categorías de FinanceBuddy).")
    dv.add("D2:D1001"); ws.add_data_validation(dv)
    dd = DataValidation(type="date", operator="greaterThan", formula1="36526", allow_blank=True, showErrorMessage=True, errorTitle="Fecha", error="Escribe una fecha, por ejemplo 28/09/2026.")
    dd.add("A2:A1001"); ws.add_data_validation(dd)
    di = DataValidation(type="decimal", operator="between", formula1="-100000000", formula2="100000000", allow_blank=True, showErrorMessage=True, errorTitle="Importe", error="Escribe un número: negativo si es un gasto, positivo si es un ingreso.")
    di.add("C2:C1001"); ws.add_data_validation(di)
    ayuda = wb.create_sheet("Cómo se usa")
    for linea in ["Plantilla de FinanceBuddy",
                  "",
                  "1. En la hoja «Movimientos» escribe una fila por movimiento: fecha, concepto, importe y categoría (se elige de la lista).",
                  "2. El importe va con signo: negativo si es un gasto (−12,50) y positivo si es un ingreso o te devuelven dinero (+30).",
                  "3. Guarda el archivo y súbelo en FinanceBuddy → Importar (o déjalo en la carpeta Importar\\Banco).",
                  "4. La primera vez te preguntará en qué cuenta va. Lo que ya estaba importado no se duplica.",
                  "",
                  "No cambies los títulos de la fila 1. La hoja «Listas» tiene tus categorías: si añades una en FinanceBuddy, descarga la plantilla de nuevo.",
                  f"Creada con FinanceBuddy {VERSION}."]:
        ayuda.append([linea])
    ayuda["A1"].font = Font(bold=True, size=14); ayuda.column_dimensions["A"].width = 120
    buf = io.BytesIO(); wb.save(buf)
    return buf.getvalue()


# ───────────── tus datos en Excel (una hoja por cosa) ─────────────
# Para guardarlos, llevarlos a otra herramienta o hacer tus propias cuentas. Es una copia para consultar: no se vuelve a importar
# tal cual (para restaurar están las copias de seguridad). Las celdas de texto nunca se interpretan como fórmulas.
_HOJAS = [
    ("Movimientos", "movimiento", [("Fecha", "fecha"), ("Tipo", "clase"), ("Concepto", "concepto"), ("Categoría", "categoria"), ("Importe", "importe"),
                                   ("Importe con signo", "_signo"), ("Cuenta", "cuenta"), ("Hacia / desde otra cuenta", "_otra"), ("Fijo enlazado", "recurrente"),
                                   ("Nota", "nota"), ("Texto del extracto", "ext_texto")]),
    ("Cuentas", "cuenta", [("Nombre", "nombre"), ("Tipo", "tipo"), ("Importo sus movimientos", "extracto"), ("IBAN (4 últimas)", "iban"), ("Notas", "notas")]),
    ("Saldos", "patrimonio", [("Fecha", "fecha"), ("Saldos por cuenta", "_saldos"), ("Valor de la inversión", "_valores"), ("Otros", "otros"), ("Deudas", "deudas"), ("Nota", "nota")]),
    ("Fijos", "recurrente", [("Nombre", "nombre"), ("Tipo", "clase"), ("Categoría", "categoria"), ("Importe", "importe"), ("Día", "dia"), ("Desde", "desde"),
                             ("Hasta", "hasta"), ("Meses", "_meses"), ("Activo", "_activo"), ("Cuenta", "cuenta"), ("Activo de inversión", "activo_inversion")]),
    ("Categorías", "categoria", [("Nombre", "nombre"), ("Grupo", "grupo"), ("Presupuesto mensual", "presupuesto"), ("Oculta", "oculta"), ("Descripción", "descripcion")]),
    ("Reglas", "regla", [("El texto contiene", "patron"), ("Tipo", "clase"), ("Categoría", "categoria"), ("Otra cuenta", "cuenta_otra"), ("Origen", "origen")]),
    ("Inversión", "activo", [("Nombre", "nombre"), ("Clase", "clase"), ("Cuenta", "cuenta"), ("ISIN", "isin"), ("Valor anotado", "valor"), ("Fecha del valor", "fecha_valor"),
                             ("Aportado al inicio", "aportado_inicial"), ("Estado", "estado"), ("Gastos anuales %", "ter")]),
    ("Aportaciones", "aportacion", [("Fecha", "fecha"), ("Activo", "activo"), ("Importe (− venta)", "importe"), ("Participaciones", "participaciones"), ("Cuenta", "cuenta"),
                                    ("Traspaso", "traspaso"), ("Ajuste", "ajuste"), ("Nota", "nota")]),
    ("Dividendos y comisiones", "cobro", [("Fecha", "fecha"), ("Activo", "activo"), ("Tipo", "tipo"), ("Importe", "importe"), ("Cuenta", "cuenta"), ("Nota", "nota")]),
    ("Objetivos", "objetivo", [("Nombre", "nombre"), ("Meta", "meta"), ("Meta en meses de gasto", "meta_meses"), ("Ahorrado", "ahorrado"), ("Cuenta", "cuenta"),
                               ("Fecha límite", "fecha_limite"), ("Prioridad", "prioridad"), ("Estado", "estado")]),
    ("Recordatorios", "recordatorio", [("Qué", "nombre"), ("Fecha", "fecha"), ("Avisar (días antes)", "avisar_dias"), ("Estado", "estado"), ("Nota", "texto")]),
]
_CLASE = {"gasto": "Gasto", "ingreso": "Ingreso", "reembolso": "Reembolso", "transferencia": "Entre cuentas", "aportacion": "Aportación"}
_FECHAS = {"fecha", "desde", "hasta", "fecha_valor", "fecha_limite"}

def _valor(tipo, r, k):
    """Valor de una celda: fechas como fecha, importes como número, texto tal cual."""
    import datetime
    if k == "_signo":
        if r.get("importe") is None or r.get("clase") == "transferencia": return None
        return -r["importe"] if r.get("clase") == "gasto" else r["importe"]
    if k == "_otra": return r.get("destino") or r.get("origen") or ""
    if k == "_saldos": return "; ".join(f"{a}: {v:.2f}" for a, v in (r.get("saldos") or {}).items())
    if k == "_valores": return "; ".join(f"{a}: {v:.2f}" for a, v in (r.get("valores") or {}).items())
    if k == "_meses": return ", ".join(str(m) for m in (r.get("meses") or [])) or "todos"
    if k == "_activo": return "no" if r.get("activo") is False else "sí"
    v = r.get(k)
    if k == "clase" and tipo in ("movimiento", "recurrente"): return _CLASE.get(v, v)
    if isinstance(v, bool): return "sí" if v else ("" if k == "oculta" else "no")
    if k in _FECHAS and v:
        try: return datetime.date.fromisoformat(v)
        except (TypeError, ValueError): return v
    return v

def datos_excel(alm):
    """Excel (bytes) con todos tus datos: una hoja por tipo (movimientos, cuentas, saldos, fijos, categorías, reglas, inversión,
    aportaciones, dividendos, objetivos, recordatorios) y una hoja «Léeme». Lo más reciente, primero."""
    import datetime
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    wb = Workbook(); wb.remove(wb.active)
    for titulo, tipo, cols in _HOJAS:
        regs = alm.todos(tipo)
        if tipo in ("movimiento", "aportacion", "cobro", "patrimonio"): regs = sorted(regs, key=lambda r: (r.get("fecha") or "", r["id"]), reverse=True)
        ws = wb.create_sheet(titulo)
        ws.append([c for c, _ in cols])
        for c in ws[1]: c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="5E8266"); c.alignment = Alignment(vertical="center", wrap_text=True)
        for r in regs:
            ws.append([_valor(tipo, r, k) for _, k in cols])
            for c in ws[ws.max_row]:
                if isinstance(c.value, str): c.data_type = "s"  # un «=1+1» o «-…» del extracto se queda como texto
                elif isinstance(c.value, datetime.date): c.number_format = "dd/mm/yyyy"
                elif isinstance(c.value, float): c.number_format = "#,##0.00"
        for i, (c, _) in enumerate(cols, 1):
            ancho = max([len(c)] + [len(str(ws.cell(f, i).value or "")) for f in range(2, min(ws.max_row, 200) + 1)])
            ws.column_dimensions[ws.cell(1, i).column_letter].width = min(48, max(11, ancho + 2))
        ws.freeze_panes = "A2"
        if regs: ws.auto_filter.ref = ws.dimensions
    ayuda = wb.create_sheet("Léeme", 0)
    for linea in ["Tus datos de FinanceBuddy",
                  "",
                  f"Exportados el {datetime.date.today():%d/%m/%Y} con FinanceBuddy {VERSION}. Es una copia para consultar o hacer tus propias cuentas: no cambia nada en la app.",
                  "En «Movimientos» el importe va sin signo y «Importe con signo» pone los gastos en negativo (las transferencias entre cuentas quedan vacías: no son ni gasto ni ingreso).",
                  "En «Aportaciones» una venta lleva el importe y las participaciones en negativo.",
                  "Para guardar algo que se pueda restaurar usa «Hacer una copia ahora» en Ajustes → Tus datos y copias.",
                  "",
                  "Hojas: " + ", ".join(t for t, _, _ in _HOJAS) + ".",
                  "",
                  "Este archivo contiene tus datos personales: guárdalo en un sitio de confianza."]:
        ayuda.append([linea])
    ayuda["A1"].font = Font(bold=True, size=14); ayuda.column_dimensions["A"].width = 130
    buf = io.BytesIO(); wb.save(buf)
    return buf.getvalue()
