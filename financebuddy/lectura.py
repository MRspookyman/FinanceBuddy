# Lectura de extractos (Excel .xlsx/.xls o CSV) y reconocimiento del formato (perfil) por la fila de cabecera.
import csv, datetime, io, os, re, unicodedata

def norm(s):
    s = unicodedata.normalize("NFD", str(s if s is not None else "").lower())
    return re.sub(r"\s+", " ", "".join(c for c in s if unicodedata.category(c) != "Mn")).strip()

def _texto_de(datos):
    """Los bytes de un archivo de texto → str. El «Texto Unicode» de Excel (UTF-16) se reconoce por su marca o por sus ceros."""
    if datos[:2] in (b"\xff\xfe", b"\xfe\xff"): return datos.decode("utf-16", "replace")
    if datos[:2000].count(b"\x00") > len(datos[:2000]) // 4:  # sin marca: uno de cada dos bytes es un cero
        return datos.decode("utf-16-le" if datos[1:2] == b"\x00" else "utf-16-be", "replace")
    for enc in ("utf-8-sig", "cp1252", "latin-1"):
        try: return datos.decode(enc)
        except UnicodeDecodeError: pass

def filas_crudas(ruta):
    """Todas las filas del archivo como listas de celdas (primera hoja si es Excel). Manda lo que el archivo es por dentro, no
    su extensión: hay bancos que llaman .xls a un .xlsx, a un texto con tabuladores o a una página web."""
    ext = os.path.splitext(ruta)[1].lower()
    with io.open(ruta, "rb") as fh: datos = fh.read()
    if datos[:4] == b"PK\x03\x04":  # un .xlsx (por dentro es un zip)
        import openpyxl
        try:
            wb = openpyxl.load_workbook(io.BytesIO(datos), data_only=True, read_only=True)
            filas = [list(r) for r in wb.worksheets[0].iter_rows(values_only=True)]
            wb.close()
        except Exception:
            raise ValueError("No se puede leer como Excel: es otro tipo de archivo o está dañado. Ábrelo y guárdalo como .xlsx o CSV.")
        return filas
    if datos[:4] == b"\xd0\xcf\x11\xe0":  # un .xls de los de antes
        import xlrd
        try: sh = xlrd.open_workbook(file_contents=datos).sheet_by_index(0)
        except Exception:
            raise ValueError("No se puede leer este Excel: puede que esté protegido con contraseña o dañado. Ábrelo y guárdalo como .xlsx o CSV.")
        out = []
        for i in range(sh.nrows):
            fila = []
            for j in range(sh.ncols):
                c = sh.cell(i, j)
                fila.append(xlrd.xldate_as_datetime(c.value, 0) if c.ctype == xlrd.XL_CELL_DATE else c.value)
            out.append(fila)
        return out
    if ext not in (".csv", ".txt", ".xls", ".xlsx", ".xlsm"): raise ValueError(f"Formato no admitido ({ext}): usa Excel (.xlsx, .xls) o CSV.")
    if not datos.strip(b"\xef\xbb\xbf\xff\xfe\x00 \t\r\n"): raise ValueError("El archivo está vacío.")
    if datos[:600].lstrip(b"\xef\xbb\xbf\xff\xfe\x00 \t\r\n")[:1] == b"<":
        raise ValueError("Este archivo no es un Excel de verdad: es una página web (o un XML) guardada con otro nombre; algunos bancos lo descargan así. "
                         "Ábrelo con Excel y guárdalo como .xlsx, o descarga el extracto en CSV.")
    texto = _texto_de(datos).replace("\r\n", "\n").replace("\r", "\n")
    muestra = "\n".join(texto.split("\n")[:40])
    sep = max([";", "\t", ","], key=lambda s: muestra.count(s))
    return list(csv.reader(io.StringIO(texto), delimiter=sep))

# Fechas que no salen con un formato fijo: con hora detrás (con o sin segundos y zona: «2026-09-30T10:15:23Z»), con el año
# de dos cifras o con el mes escrito («30 sep 2026», «30 de septiembre de 2026», «Sep 30, 2026»). Siempre la celda entera.
# No se intenta el orden americano (09/30/2026) ni «20260930»: son ambiguos y un número cualquiera pasaría por fecha.
_HORA = r"(?:[ t,]+\d{1,2}:\d{2}(?::\d{2}(?:[.,]\d+)?)?\s*(?:z|utc|gmt|[+-]\d{2}(?::?\d{2})?)?)?"
RE_FECHAS = [(re.compile(p + _HORA), orden) for p, orden in (  # orden: en qué grupo están el año, el mes y el día
    (r"(\d{4})-(\d{1,2})-(\d{1,2})", (1, 2, 3)),
    (r"(\d{1,2})[/.\-](\d{1,2})[/.\-](\d{4}|\d{2})", (3, 2, 1)),
    (r"(\d{1,2})[ /.\-]+(?:de )?([a-z]{3,10})\.?[ /.\-,]+(?:de )?(\d{4})", (3, 2, 1)),
    (r"([a-z]{3,10})\.? (\d{1,2}),? (\d{4})", (3, 1, 2)))]
NOMBRES_MES = ["enero", "febrero", "marzo", "abril", "mayo", "junio", "julio", "agosto", "septiembre", "octubre", "noviembre", "diciembre",
               "january", "february", "march", "april", "may", "june", "july", "august", "september", "october", "november", "december"]

def _mes(t):
    """«9», «sep», «sept», «septiembre», «september» → 9. None si no es un mes («martes» no es marzo)."""
    if t.isdigit(): return int(t)
    if t in ("set", "setiembre"): t = "septiembre"
    return next((i % 12 + 1 for i, n in enumerate(NOMBRES_MES) if n.startswith(t)), None)

def fecha(v):
    if isinstance(v, datetime.datetime): return v.date().isoformat()
    if isinstance(v, datetime.date): return v.isoformat()
    if isinstance(v, (int, float)) and 32874 <= v < 73051:  # número de serie de Excel (de 1990 a 2099)
        return (datetime.date(1899, 12, 30) + datetime.timedelta(days=int(v))).isoformat()
    s = str(v or "").strip()
    for f in ("%d/%m/%Y", "%d-%m-%Y", "%Y-%m-%d", "%d/%m/%y", "%d-%m-%y", "%d.%m.%Y", "%Y/%m/%d", "%Y-%m-%d %H:%M:%S", "%d/%m/%Y %H:%M:%S", "%d/%m/%Y %H:%M"):
        try: return datetime.datetime.strptime(s, f).date().isoformat()
        except ValueError: pass
    t = norm(s)
    for patron, (a, m, d) in RE_FECHAS:
        x = patron.fullmatch(t)
        if not x: continue
        año, mes = int(x.group(a)), _mes(x.group(m))
        if año < 100: año += 2000 if año < 69 else 1900
        try: return datetime.date(año, mes, int(x.group(d))).isoformat()
        except (TypeError, ValueError): return None
    return None

def con_cifras(v):
    """¿La celda trae algo con cifras? Una fila con cifras donde va la fecha y donde va el importe era un movimiento; una línea
    en blanco, un rótulo o un total sin fecha, no."""
    return v is not None and not isinstance(v, bool) and bool(re.search(r"\d", str(v)))

def numero(v, miles=True):
    """1.234,56 · -12,30 · 1,234.56 · 1234.5 · «12,30 €» → float. None si no es un número.
    «1.234» (un solo punto, sin coma y tres cifras detrás) es 1234, con el punto de miles español; con miles=False (columnas
    de participaciones, o un archivo que usa el punto como coma decimal: ver punto_decimal) es 1,234."""
    if isinstance(v, bool) or v is None: return None
    if isinstance(v, (int, float)): return float(v)
    s = str(v).strip().replace("€", "").replace("EUR", "").replace("\xa0", "").replace(" ", "")
    s = s.replace("−", "-").replace("–", "-")  # el menos tipográfico («−») y la raya corta, al copiar de la web del banco
    neg = s.startswith("(") and s.endswith(")")
    s = s.strip("()")
    if not s or not re.fullmatch(r"[+-]?[\d.,]+-?", s): return None
    if s.endswith("-"): s = "-" + s[:-1]
    if "," in s and "." in s:
        s = s.replace(".", "").replace(",", ".") if s.rfind(",") > s.rfind(".") else s.replace(",", "")
    elif "," in s: s = s.replace(",", ".") if s.count(",") == 1 else s.replace(",", "")
    elif s.count(".") > 1: s = s.replace(".", "")  # 1.234.567
    elif miles and re.fullmatch(r"[+-]?[1-9]\d{0,2}\.\d{3}", s): s = s.replace(".", "")  # 1.234 (otro punto solo, decimal: 12.5)
    try: n = float(s)
    except ValueError: return None
    return -n if neg else n

def punto_decimal(valores):
    """¿Las cifras usan el punto como coma decimal («12.50», «1234.5»)? Entonces un «1.234» de esas mismas columnas es 1,234
    y no mil doscientos treinta y cuatro. Se decide por archivo, mirando todos los valores (no fila a fila)."""
    return any(isinstance(v, str) and "," not in v and v.count(".") == 1 and numero(v, miles=False) is not None
               and numero(v) == numero(v, miles=False) for v in valores)

def texto(v):
    if v is None: return ""
    if isinstance(v, float) and v.is_integer(): v = int(v)
    return re.sub(r"\s+", " ", str(v)).strip()

# ───────────── cabecera y perfiles ─────────────
def buscar_cabecera(filas, columnas):
    """Índice de la primera fila (de las 40 primeras) que contiene todos los textos de `columnas` (normalizados)
    y el índice de cada columna. None si no está."""
    buscados = {k: norm(v) for k, v in columnas.items() if v}
    for i, f in enumerate(filas[:40]):
        cab = [norm(x) for x in f]
        idx = {}
        for k, v in buscados.items():
            j = next((j for j, c in enumerate(cab) if c == v), None)
            if j is None: j = next((j for j, c in enumerate(cab) if c.startswith(v)), None)
            if j is None: break
            idx[k] = j
        else:
            if len(set(idx.values())) == len(idx): return i, idx
    return None

def reconocer(filas, perfiles, tipo=None):
    """Primer perfil (del tipo pedido, si se pide) cuya cabecera aparece en el archivo → (perfil, i_cabecera, índices).
    Los perfiles con más columnas se prueban antes (son más concretos)."""
    cand = [p for p in perfiles if not tipo or p.get("tipo") == tipo]
    for p in sorted(cand, key=lambda p: -len(p.get("columnas") or {})):
        r = buscar_cabecera(filas, p.get("columnas") or {})
        if r: return p, r[0], r[1]
    return None

RE_IBAN = re.compile(r"\b([A-Z]{2}\d{2}(?:\s*[\dA-Z]){10,30})\b")

def datos_cabecera(filas, hasta=40):
    """Lo que dicen las filas de encima de la tabla (si lo dicen): IBAN (solo sus 4 últimas cifras) y titular.
    Santander: «Cuenta | Fecha» y debajo el IBAN; «Titular | Saldo» y debajo el nombre."""
    out = {}
    filas = filas[:hasta]
    for i, f in enumerate(filas):
        for j, c in enumerate(f):
            t = texto(c)
            if not t: continue
            m = RE_IBAN.search(t.upper())
            if m and "iban" not in out:
                digitos = re.sub(r"\D", "", m.group(1)[4:])
                if len(digitos) >= 10: out["iban"] = digitos[-4:]
            if norm(t).rstrip(":") in ("titular", "titulares", "nombre del titular", "titular de la cuenta") and "titular" not in out:
                cand = [filas[i + 1][j] if i + 1 < len(filas) and j < len(filas[i + 1]) else None, f[j + 1] if j + 1 < len(f) else None]
                etiqueta = lambda x: norm(x).rstrip(":") in ("saldo", "fecha", "cuenta", "iban", "divisa", "moneda", "entidad", "oficina")
                nombre = next((texto(x) for x in cand if texto(x) and not etiqueta(texto(x)) and not RE_IBAN.search(texto(x).upper())
                               and numero(x) is None and fecha(x) is None), None)
                if nombre and len(nombre) >= 5: out["titular"] = nombre
    return out

def cabecera_probable(filas):
    """Para configurar un formato nuevo: la fila que parece la cabecera (≥ 3 textos y debajo una fila con fecha e importe)."""
    for i, f in enumerate(filas[:40]):
        textos = [x for x in f if isinstance(x, str) and x.strip() and fecha(x) is None and numero(x) is None]
        if len(textos) < 3: continue
        sig = filas[i + 1] if i + 1 < len(filas) else []
        if any(fecha(x) for x in sig) and any(numero(x) is not None and not fecha(x) for x in sig): return i
    return 0

# Cómo llama cada banco a cada columna (lo que se ha visto en extractos españoles). Se compara con la cabecera normalizada:
# primero entera, luego «empieza por» y por último «contiene». El orden de cada lista va de lo más claro a lo más genérico.
NOMBRES = {
    "fecha": ["fecha operacion", "f. operacion", "f.operacion", "fecha de operacion", "fecha contable", "fecha movimiento", "fecha de la operacion",
              "started date", "completed date", "transaction date", "fecha", "date"],
    "fecha_valor": ["fecha valor", "f. valor", "f.valor", "fecha de valor", "value date", "valor"],
    "concepto": ["concepto", "descripcion", "description", "detalle", "concepto del movimiento", "descripcion del movimiento", "movimiento",
                 "operacion", "referencia", "beneficiario", "comercio", "merchant", "texto"],
    "importe": ["importe", "importe eur", "importe (eur)", "importe €", "cantidad", "monto", "importe del movimiento", "amount", "euros"],
    "cargo": ["cargo", "debe", "debito", "salida", "paid out", "money out", "pagos", "pago", "gasto"],
    "abono": ["abono", "haber", "credito", "entrada", "paid in", "money in", "ingresos", "ingreso", "cobro"],
    "saldo": ["saldo", "saldo eur", "saldo (eur)", "saldo posterior", "saldo disponible", "saldo contable", "balance"],
    "categoria": ["categoria", "categoría", "tipo de gasto", "category"],
}

def _columna_de(campo, cab):
    """Índice de la columna que parece ser `campo` según su nombre (o None)."""
    for nombre in NOMBRES[campo]:
        for prueba in (lambda c, n: c == n, lambda c, n: c.startswith(n), lambda c, n: n in c):
            j = next((j for j, c in enumerate(cab) if c and prueba(c, nombre)), None)
            if j is not None: return j
    return None

def columnas_probables(cabecera, ejemplos):
    """Qué columna es cada cosa en un archivo de un banco que la app no conoce, mirando los nombres de la cabecera y
    comprobándolo con las filas de ejemplo (que la de fecha traiga fechas y la de importe, números). → {campo: texto de la cabecera}.
    Lo que no se vea claro se deja fuera: el usuario lo elige en la pantalla."""
    cab = [norm(c) for c in cabecera]
    def valores(j):
        return [f[j] for f in ejemplos if j < len(f) and texto(f[j])]
    def es_fecha(j, por_nombre=True):
        # Una celda de número (45930) puede ser una fecha de Excel o un importe: por el contenido solo, hace falta alguna fecha de verdad
        v = valores(j)
        return bool(v) and all(fecha(x) for x in v) and (por_nombre or any(not isinstance(x, (int, float)) for x in v))
    def es_numero(j, vacia_vale=False):
        v = valores(j)
        if not v: return vacia_vale  # en «Debe/Haber» cada fila rellena solo una de las dos columnas
        return all(numero(x) is not None for x in v)
    out, usadas = {}, set()
    for campo in ("fecha", "fecha_valor", "saldo", "importe", "cargo", "abono", "categoria", "concepto"):
        j = _columna_de(campo, cab)
        if j is None or j in usadas or not texto(cabecera[j]): continue
        if campo in ("fecha", "fecha_valor") and ejemplos and not es_fecha(j): continue
        if campo in ("importe", "saldo") and ejemplos and not es_numero(j): continue
        if campo in ("cargo", "abono") and ejemplos and not es_numero(j, vacia_vale=True): continue
        if campo == "concepto" and ejemplos and (es_fecha(j) or es_numero(j)): continue
        out[campo] = texto(cabecera[j]); usadas.add(j)
    if "importe" in out: out.pop("cargo", None); out.pop("abono", None)  # con el importe con signo no hacen falta
    elif not ("cargo" in out and "abono" in out): out.pop("cargo", None); out.pop("abono", None)
    if "fecha" not in out and "fecha_valor" in out: out["fecha"] = out.pop("fecha_valor")  # solo hay fecha valor: esa es la fecha
    # Sin nombres reconocibles, por el contenido: la primera columna con fechas y la primera con números distintos del saldo
    if "fecha" not in out:
        j = next((j for j in range(len(cab)) if j not in usadas and es_fecha(j, por_nombre=False)), None)
        if j is not None: out["fecha"] = texto(cabecera[j]); usadas.add(j)
    if "importe" not in out and not ("cargo" in out and "abono" in out):
        j = next((j for j in range(len(cab)) if j not in usadas and es_numero(j)), None)
        if j is not None: out["importe"] = texto(cabecera[j]); usadas.add(j)
    if "concepto" not in out:
        j = next((j for j in range(len(cab)) if j not in usadas and texto(cabecera[j]) and not es_fecha(j) and not es_numero(j)), None)
        if j is not None: out["concepto"] = texto(cabecera[j])
    return out if "fecha" in out and "concepto" in out and ("importe" in out or ("cargo" in out and "abono" in out)) else {}

def muestra_para_configurar(filas):
    i = cabecera_probable(filas)
    cab = [texto(x) for x in filas[i]] if filas else []
    ejemplos = [[texto(x.isoformat()[:10] if isinstance(x, (datetime.date, datetime.datetime)) else x) for x in f] for f in filas[i + 1:i + 6]]
    return {"fila_cabecera": i, "cabecera": cab, "ejemplos": ejemplos, "columnas_probables": columnas_probables(cab, ejemplos)}
